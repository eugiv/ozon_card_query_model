import copy
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from tools.defense import (
    ANALYSIS,
    REPOSITORY,
    SUMMARY,
    TOTAL,
    build,
    refresh_summary,
    results_summary,
    save_materials,
)


@pytest.fixture
def facts():
    return json.loads(SUMMARY.read_text()), json.loads(ANALYSIS.read_text())


def all_text(p):
    return "\n".join(shape.text for slide in p.slides for shape in slide.shapes if shape.has_text_frame)


def test_complete_defense_story_and_sources(facts):
    p, _ = build(*facts)
    assert len(p.slides) == TOTAL
    text = all_text(p)
    for term in ("[ФИО автора]", "ПОСТАНОВКА ЗАДАЧИ", "АНАЛИЗ", "МЕТОДИКА РЕШЕНИЯ",
                 "РЕЗУЛЬТАТЫ", "NDCG", "22,85%", "3 530", "353 товаров", "PairLogit",
                 "сентябрь", "материала", "меламиновая губка", "LightGCN", "InfoNCE",
                 "46 787", "23 569", "10 742", "17,95%", "3 203 109", "600", REPOSITORY):
        assert term.lower() in text.lower()
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in p.slides)
    assert "Transformer" in notes and "замороженные веса" in notes
    assert "не независимый test" in notes
    assert "Фото" in notes


def test_metrics_come_from_snapshot_and_inputs_are_not_modified(facts):
    source = copy.deepcopy(facts)
    summary = copy.deepcopy(facts[0])
    summary["test"]["models"]["learned_fusion"]["recall@10"] = .2134
    p, _ = build(summary, facts[1])
    assert "21,34%" in all_text(p)
    assert facts == source


def test_model_status_distinguishes_pretraining_training_and_selection(facts):
    p, _ = build(*facts)
    search_text = " ".join(" ".join(shape.text.split()) for shape in p.slides[10].shapes if shape.has_text_frame)
    training_text = " ".join(shape.text for shape in p.slides[12].shapes if shape.has_text_frame)
    assert "Предобученная E5" in search_text
    assert "используем её обученные веса" in search_text
    assert "Три модели обучены на данных проекта" in training_text
    for model in ("CatBoostRanker", "Адаптер E5", "LightGCN"):
        assert model in training_text
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in p.slides)
    assert "используются без изменения" in notes
    assert "обученный альтернативный ранжировщик" in notes
    assert "E5, co-sales, адаптер и LightGCN по 25%" in notes
    for misleading in ("Transformer не дообучен", "не обученный LightGCN", "не выбран итоговым методом",
                       "Сам Transformer и cross-encoder не дообучены"):
        assert misleading not in notes + all_text(p)


def test_presentation_has_visual_hierarchy_and_comparison_chart(facts):
    p, _ = build(*facts)
    assert str(p.slides[0].background.fill.fore_color.rgb) == "142B35"
    assert str(p.slides[1].background.fill.fore_color.rgb) == "FFFFFF"
    assert any(shape.has_chart and shape.chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED
               for shape in p.slides[14].shapes)
    conclusion = " ".join(shape.text for shape in p.slides[20].shapes if shape.has_text_frame)
    assert "Задача решена" in conclusion
    assert "Практическая ценность" in conclusion
    assert "SEO-специалиста" in conclusion
    assert "прирост продаж" in p.slides[20].notes_slide.notes_text_frame.text


def test_source_dates_and_quality_flags_are_not_baked_into_slide_copy(facts):
    summary = copy.deepcopy(facts[0])
    summary["source"]["first_day"] = "2026-02-03"
    summary["source"]["last_day"] = "2026-10-09"
    summary["source"]["missing_days"] = ["2026-03-04"]
    summary["config"]["test_decision"] = "2026-09-20"
    p, _ = build(summary, facts[1])
    text = all_text(p)
    assert "03.02–09.10" in text
    assert "Дней без записей: 1 · 04.03" in text
    assert "21.09–04.10" in text


def test_mismatched_source_or_evaluation_requires_review(facts):
    summary = copy.deepcopy(facts[0])
    summary["source"]["fact_rows"] -= 1
    with pytest.raises(ValueError, match="расходятся"):
        build(summary, facts[1])
    summary = copy.deepcopy(facts[0])
    summary["evaluation_status"] = "independent_test"
    with pytest.raises(ValueError, match="статуса"):
        build(summary, facts[1])


def test_readable_fonts_editable_charts_and_real_hyperlink(facts):
    p, _ = build(*facts)
    charts = [shape.chart for slide in p.slides for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 7
    for slide in p.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        assert run.font.size.pt >= 14
                        if shape.top > 2286000 and shape.top < 9601200:
                            assert run.font.size.pt >= 18
    links = [run.hyperlink.address for slide in p.slides for shape in slide.shapes if shape.has_text_frame
             for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.hyperlink.address]
    assert REPOSITORY in links
    values = list(charts[3].series[0].values)
    assert values[-1] == facts[0]["test"]["models"]["learned_fusion"]["recall@10"]


def test_public_snapshot_is_reproducible_from_actual_saved_aggregates(tmp_path, facts, monkeypatch):
    """Синтетические файлы по публичной схеме; приватные модели не требуются."""
    summary, analysis = facts
    analysis_path = tmp_path / "analysis.json"
    analysis_path.write_text(json.dumps(analysis))
    monkeypatch.setattr("tools.defense.ANALYSIS", analysis_path)
    release = {
        **summary["current_queue"], "as_of": summary["current_queue"]["as_of"],
        "source_end": summary["source"]["last_day"],
        "review_pairs": summary["current_queue"]["recommended_review_pairs"],
        "training_rows": summary["final_training"]["rows"],
        "training_skus": summary["final_training"]["skus"],
        "training_groups": summary["final_training"]["groups"],
        "limitations": ["synthetic fixture"],
    }

    def models(part):
        counts = {k: part[k] for k in ("evaluated_skus", "positive_pairs", "candidate_pairs", "retrieval_recall")}
        return {k: {**counts, **v} for k, v in part["models"].items()}

    reports = {
        "baseline": {"config": summary["config"], "dictionary": summary["dictionary"],
                     "training": summary["training"], "selected_iterations": summary["selected_iterations"]},
        "evaluation": {"source": summary["source"], "aggregation": summary["aggregation"],
                       "selected_method": summary["selected_method"], "fusion": summary["fusion"],
                       "test": {"models": models(summary["test"]), "zero_views": models(summary["zero_views_test"])},
                       "new_days": {**summary["new_days"], "models": models(summary["new_days"])},
                       "limitations": ["synthetic fixture"]},
        "release": release, "ablation": summary["ablation"], **summary["neural_training"],
    }
    config = tmp_path / "config.json"
    config.write_text(json.dumps({"outputs": {key: key for key in reports}}))
    for key, value in reports.items():
        (tmp_path / key).mkdir()
        (tmp_path / key / "metrics.json").write_text(json.dumps(value))
    output = tmp_path / "public.json"
    result = refresh_summary(config, output)
    for key in ("source", "test", "zero_views_test", "ablation", "training", "final_training", "current_queue"):
        assert result[key] == summary[key]
    assert json.loads(output.read_text()) == result
    assert output.stat().st_size < 20_000


def test_current_methods_and_boundaries_are_explicit_on_slides(facts):
    p, _ = build(*facts)
    text = all_text(p)
    for outdated in ("18,23%", "3 610", "361 товаров", "50% E5", "только планируются"):
        assert outdated not in text
    for term in ("одинаковом", "финальное обучение", "без обучения", "4 дня",
                 "Включает ноль", "Фото — вне текущего", "9 ML-модулей"):
        assert term.lower() in text.lower()
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in p.slides)
    assert "37 245" in notes
    assert "04.09" in notes and "07.10" in notes
    assert "не независимый test" in notes


def test_saved_presentation_theme_and_package(tmp_path):
    output = tmp_path / "defense.pptx"
    save_materials(output, tmp_path / "contact.png")
    assert len(Presentation(output).slides) == TOTAL
    with ZipFile(output) as archive:
        assert archive.testzip() is None
        theme = archive.read("ppt/theme/theme1.xml").decode()
        assert 'name="Тема"' in theme
        assert 'val="44CCA0"' in theme
        assert len([name for name in archive.namelist() if name.startswith("ppt/embeddings/")]) == 7
    assert (tmp_path / "contact.png").is_file()
    assert not list(Path(tmp_path).glob("defense-*/"))


def test_wrong_output_cannot_overwrite_source(tmp_path):
    output = tmp_path / "source.md"
    output.write_text("сохранить")
    with pytest.raises(ValueError, match="расширения"):
        save_materials(output)
    assert output.read_text() == "сохранить"


def test_results_include_actual_stress_and_loss_diagnosis(facts):
    p, _ = build(*facts)
    text = all_text(p)
    for term in ("ШУМ И ВЫБРОСЫ", "Потеря", "1%", "Не найдены поиском", "разные способы усреднения"):
        assert term in text
    charts = [shape.chart for slide in p.slides for shape in slide.shapes if shape.has_chart]
    stage = facts[0]["stage4"]["test"]
    keys = ["clean", "score_noise_05", "score_noise_15", "counter_outliers", "candidate_loss_10", "candidate_loss_30"]
    assert list(charts[-1].series[0].values) == [stage["scenarios"][key]["recall@10"] for key in keys]
    distribution = stage["rank_distribution"]
    assert sum(value["macro_share"] for value in distribution.values()) == pytest.approx(1)
    assert sum(value["positive_pairs"] for value in distribution.values()) == facts[0]["test"]["positive_pairs"]
    assert distribution["missing_from_pool"]["macro_share"] == pytest.approx(1 - facts[0]["test"]["retrieval_recall"])
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in p.slides)
    assert "не всей системы поиска" in text
    assert "не опечатки" in notes


@pytest.mark.parametrize("damage", [None, "checksum", "parent", "method", "metrics", "counts"])
def test_results_summary_validates_provenance_and_only_exports_aggregates(tmp_path, facts, damage):
    summary, _ = facts
    public = summary["stage4"]
    source = {key: public[key] for key in ("status", "seed", "repetitions", "limitations")}
    source.update(selected_method=summary["selected_method"],
                  fusion={key: summary["fusion"][key] for key in ("weights", "offset")})
    buckets = {"top_5": "1-5", "rank_6_10": "6-10", "rank_11_50": "11-50",
               "rank_over_50": ">50", "missing_from_pool": "missing_from_pool"}
    for split in ("test", "new_days"):
        cohort = summary[split]
        counts = {key: cohort[key] for key in ("evaluated_skus", "positive_pairs", "candidate_pairs")}
        distribution = public[split]["rank_distribution"]
        source[split] = {
            "horizon_days": public[split]["horizon_days"],
            "clean_comparisons": {method: {**counts, **metrics} for method, metrics in cohort["models"].items()},
            "rank_distribution": {
                "macro_fractions": {key: distribution[name]["macro_share"] for name, key in buckets.items()},
                "pair_counts": {key: distribution[name]["positive_pairs"] for name, key in buckets.items()},
            },
            "scenarios": {},
        }
        for key, scenario in public[split]["scenarios"].items():
            source[split]["scenarios"][key] = {
                "models": {"selected": {
                    "recall@10": {"mean": scenario["recall@10"], "std": scenario["recall@10_std"]},
                    "ndcg@10": {"mean": scenario["ndcg@10"]},
                }},
                "selected_top10": {"intersection_over_k": {"mean": scenario["overlap@10"]}},
                "selected_delta_vs_clean": scenario["delta_recall@10"],
            }
    if damage == "method":
        source["selected_method"] = "popularity"
    elif damage == "metrics":
        source["test"]["clean_comparisons"]["lexical"]["recall@10"] += .1
    elif damage == "counts":
        source["test"]["clean_comparisons"]["lexical"]["positive_pairs"] += 1
    evaluation = tmp_path / "evaluation"
    evaluation.mkdir()
    (evaluation / "manifest.json").write_text("synthetic fixture, not a real experiment")
    metrics = tmp_path / "metrics.json"
    metrics.write_text(json.dumps(source))
    digest = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = {"files": {"metrics.json": digest(metrics)},
                "parents": {str(evaluation): digest(evaluation / "manifest.json")}}
    if damage == "checksum":
        manifest["files"]["metrics.json"] = "0" * 64
    elif damage == "parent":
        manifest["parents"][str(evaluation)] = "0" * 64
    (tmp_path / "manifest.json").write_text(json.dumps(manifest))
    if damage:
        with pytest.raises(ValueError):
            results_summary(metrics, evaluation, summary)
    else:
        assert results_summary(metrics, evaluation, summary) == public


def test_all_six_sections_and_concrete_conclusions_are_visible(facts):
    p, _ = build(*facts)
    text = all_text(p)
    for chapter in ("5 · ВЫВОДЫ И ЗАКЛЮЧЕНИЕ", "5 · ВЫВОДЫ / НАПРАВЛЕНИЯ РАЗВИТИЯ",
                    "6 · СПИСОК ИСТОЧНИКОВ / МЕТОДЫ", "6 · СПИСОК ИСТОЧНИКОВ / РЕАЛИЗАЦИЯ"):
        assert chapter in text
    route_text = " ".join(shape.text for shape in p.slides[1].shapes if shape.has_text_frame)
    for section in ("Постановка задачи", "Анализ данных", "Методика решения", "Результаты",
                    "Выводы и заключение", "Список источников"):
        assert section in route_text
    conclusion = " ".join(shape.text for shape in p.slides[20].shapes if shape.has_text_frame)
    for fragment in ("19,06%", "22,85%", "+3,79 п.п.", "без повторного обучения",
                     "Практическая ценность", "Отсев дублей", "3 530"):
        assert fragment in conclusion
    notes = p.slides[20].notes_slide.notes_text_frame.text
    assert "Сентябрь уже просмотрен" in notes
    assert "A/B и рост продаж — последующая бизнес-задача" in notes
    roadmap = " ".join(shape.text for shape in p.slides[21].shapes if shape.has_text_frame)
    for fragment in ("26,87%", "50,28%", "14 дней", "Полнота пула", "NDCG",
                     "Соответствие товару", "неизвестно", "Фото — вне текущего объёма"):
        assert fragment in roadmap


def test_conclusion_numbers_are_from_aggregate_snapshot(facts):
    summary = copy.deepcopy(facts[0])
    summary["ablation"]["test"]["e5_co_sales"]["recall@10"] = .18
    summary["ablation"]["test"]["all_four"]["recall@10"] = .24
    summary["current_queue"]["recommended_review_pairs"] = 3210
    p, _ = build(summary, facts[1])
    conclusion = " ".join(shape.text for shape in p.slides[20].shapes if shape.has_text_frame)
    for fragment in ("18,00%", "24,00%", "+6,00 п.п.", "3 210"):
        assert fragment in conclusion


def test_bibliography_matches_note_and_separates_methods_from_our_results(facts):
    p, _ = build(*facts)
    sources = list(p.slides)[-2:]
    urls = {run.hyperlink.address for slide in sources for shape in slide.shapes if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.hyperlink.address}
    note = (SUMMARY.parents[1] / "docs/explanatory_note.md").read_text()
    note_urls = set(re.findall(r"\]\((https?://[^)]+)\)", note)) - {REPOSITORY}
    assert len(urls) == 12
    assert urls == note_urls
    text = " ".join(shape.text for slide in sources for shape in slide.shapes if shape.has_text_frame)
    for fragment in ("BPR", "InfoNCE", "PairLogit", "multilingual-e5-small",
                     "Статьи обосновывают методы", "контрольными суммами"):
        assert fragment in text