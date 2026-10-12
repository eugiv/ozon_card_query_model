import json
import math
import re

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from tools import defense as note_module
from tools.defense import (
    DEFAULT_SOURCE,
    LINK,
    PAGEBREAK,
    create_explanatory_note,
    parse_pages,
)


@pytest.fixture(scope="module")
def note(tmp_path_factory):
    folder = tmp_path_factory.mktemp("explanatory_note")
    output = create_explanatory_note(output=folder / "note.docx")
    return Document(str(output))


def test_ten_logical_pages_and_contents(note):
    pages = parse_pages(DEFAULT_SOURCE.read_text(encoding="utf-8"))
    assert len(pages) == 10
    breaks = note.element.body.xpath('.//w:br[@w:type="page"]')
    assert len(breaks) == 9
    headings = [p.text for p in note.paragraphs if p.style.name in ("Title", "Heading 1")]
    assert headings == [page[0][2:] for page in pages]
    contents = note.tables[0]
    assert len(contents.rows) == 11
    for number, row in enumerate(contents.rows[1:], 1):
        assert row.cells[1].text == str(number)
        assert row.cells[0].text == ("Титульный лист" if number == 1 else headings[number - 1])
    assert abs(contents.columns[1].width - Mm(25)) < Mm(0.1)


def test_source_and_docx_have_identical_visible_text(note):
    expected = []
    for page in parse_pages(DEFAULT_SOURCE.read_text(encoding="utf-8")):
        for block in page:
            if block.startswith("|"):
                for line in block.splitlines():
                    cells = [cell.strip() for cell in line.strip("|").split("|")]
                    if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                        continue
                    expected.extend(LINK.sub(r"\1", cell) for cell in cells)
            else:
                expected.append(LINK.sub(r"\1", block.removeprefix("# ")))
    actual = " ".join(
        "".join(node.text or "" for node in paragraph.iter(qn("w:t")))
        for paragraph in note.element.body.iter(qn("w:p"))
    )
    visible_expected = re.sub(r"`([^`]+)`", r"\1", " ".join(expected))
    assert actual.split() == visible_expected.split()


def test_format_fields_and_editable_links(note):
    section = note.sections[0]
    assert abs(section.page_width - Mm(210)) < Mm(0.1)
    assert abs(section.page_height - Mm(297)) < Mm(0.1)
    for side in ("top", "bottom", "left", "right"):
        assert abs(getattr(section, f"{side}_margin") - Mm(20)) < Mm(0.1)
    normal = note.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == pytest.approx(1.15)
    for name in ("Title", "Heading 1"):
        assert note.styles[name].font.name == "Times New Roman"
        assert note.styles[name].font.size == Pt(12)
        assert note.styles[name].paragraph_format.line_spacing == pytest.approx(1.15)
    assert "Малотрафиковые" in section.header.paragraphs[0].text
    assert section.footer._element.xpath('.//w:fldSimple[@w:instr="PAGE"]')
    urls = {r.target_ref for r in note.part.rels.values()
            if r.reltype == RELATIONSHIP_TYPE.HYPERLINK and r.is_external}
    assert "https://github.com/eugiv/ozon_card_query_model" in urls
    assert "https://arxiv.org/abs/1706.09516" in urls
    assert "https://arxiv.org/abs/2002.02126" in urls
    assert "https://arxiv.org/abs/1205.2618" in urls
    assert "https://arxiv.org/abs/1807.03748" in urls
    assert "https://huggingface.co/intfloat/multilingual-e5-small" in urls
    assert len(urls) == 13
    assert not note.element.body.xpath(".//w:drawing")


def test_metrics_and_honest_status():
    text = DEFAULT_SOURCE.read_text(encoding="utf-8")
    # Выбранные метрики этапа 3; приватные модели и выгрузки тесту не нужны.
    snapshot = json.loads((DEFAULT_SOURCE.parents[1] / "reports/experiment_summary.json").read_text())
    rows = (
        ("TF-IDF", "lexical"), ("Co-sales", "co_sales"),
        ("Предобученная E5", "pretrained_embeddings"), ("Обученный адаптер", "semantic_adapter"),
        ("LightGCN", "lightgcn"), ("CatBoost", "catboost_pu"), ("Гибрид четырёх", "learned_fusion"),
    )
    for name, method in rows:
        metrics = snapshot["test"]["models"][method]
        recall, precision, ndcg = (metrics[key] for key in ("recall@10", "precision@10", "ndcg@10"))
        expected = f"| {name} | {recall * 100:.2f}% | {precision * 100:.2f}% | {ndcg:.4f} |"
        assert expected.replace(".", ",") in text
    for number in (3203109, 2171, 268, 477579, 191456, 232791, 37245, 334, 2504,
                   46787, 3053, 413, 121, 260, 166291, 3530, 353, 212787, 27055, 907,
                   94, 180, 96, 132, 23569, 589, 9525, 10742):
        assert f"{number:,}".replace(",", " ") in text
    for value in (0.8224372935379805, 0.731267217630854, 0.1905844155844156,
                  0.21592876820149548, 0.2089662862390135, 0.15138888888888888,
                  0.1795138888888889):
        assert f"{value * 100:.2f}%".replace(".", ",") in text
    for fragment in ("[ФИО автора]", "[Образовательная программа]", "[ФИО руководителя]",
                     "просмотры, не клики", "test уже просмотрен", "25% E5",
                     "25% co-sales", "25% адаптер", "25% LightGCN",
                     "not_verified", "add_candidate не создаётся",
                     "перекрывающиеся ответы не являются независимыми",
                     "+0,20…+7,57", "−2,08…+7,92", "интервал включает ноль",
                     "384 → 128 → 384", "InfoNCE", "BPR", "37 245", "2 504",
                     "600 итераций", "CatBoost", "обученный альтернативный ранжировщик",
                     "Фото и OCR в текущем решении не использовались",
                     "Публичная копия не содержит .env, приватных данных и весов"):
        assert fragment in text
    assert "нет" in text.split("измеренного бизнес-прироста", 1)[1].split(".", 1)[0]
    for retired in ("full_history_20261007.md", "experiment_summary.json", "project_map",
                    "submission", "manual_label"):
        assert retired not in text
    assert "src/recommender.py run --config config.json" in text
    assert "src/recommender.py verify --config config.json" in text
    assert "src/recommender.py recommend --release reports/stage3_release" in text
    assert "--as-of 2026-10-11T01:28:00Z --output reports/recommendations_new" in text
    assert "He X." in text and "2020" in text
    assert "Rendle S." in text and "BPR" in text and "2009" in text


def test_inputs_not_modified_and_invalid_source_rejected(tmp_path):
    before = DEFAULT_SOURCE.read_bytes()
    create_explanatory_note(output=tmp_path / "note.docx")
    assert DEFAULT_SOURCE.read_bytes() == before
    with pytest.raises(ValueError, match="отдельным"):
        create_explanatory_note(DEFAULT_SOURCE, DEFAULT_SOURCE)
    with pytest.raises(ValueError, match="десять"):
        parse_pages("# Одна страница")
    with pytest.raises(ValueError, match="бюджет"):
        parse_pages(PAGEBREAK.join(["# Заголовок\n\n" + "слово " * 401] * 10))


def test_results_stress_numbers_match_public_aggregate_snapshot():
    text = DEFAULT_SOURCE.read_text(encoding="utf-8")
    snapshot = json.loads((DEFAULT_SOURCE.parents[1] / "reports/experiment_summary.json").read_text())
    stage = snapshot["stage4"]["test"]
    for scenario in stage["scenarios"].values():
        value = f"{scenario['recall@10'] * 100:.2f}%".replace(".", ",")
        assert value in text
    for bucket in ("rank_11_50", "rank_over_50", "missing_from_pool"):
        value = f"{stage['rank_distribution'][bucket]['macro_share'] * 100:.2f}%".replace(".", ",")
        assert value in text
    for phrase in ("без обучения и подбора весов", "seed 7/8/9", "×100 у 1%",
                   "не опечатки", "Не доказывает устойчивость всей системы", "src/recommender.py results"):
        assert phrase in text


def test_conclusion_connects_achievements_limits_and_development_to_evidence(note):
    pages = parse_pages(DEFAULT_SOURCE.read_text(encoding="utf-8"))
    conclusion = " ".join(pages[8])
    for phrase in ("работоспособная воспроизводимая гибридная", "3 530", "19,06%", "22,85%",
                   "+3,79 п.п.", "+19,89%", "Что измеряет эксперимент",
                   "26,87%", "зафиксировать версию", "Полнота пула", "NDCG",
                   "конфликты и неизвестные отдельно", "бизнес-решения"):
        assert phrase.lower() in conclusion.lower()
    scope = " ".join(pages[4])
    for phrase in ("Что измеряет эксперимент", "test уже просмотрен", "не независимое подтверждение",
                   "интервал включает ноль", "фото/OCR вне текущего объёма", "95,04%",
                   "не объявляются неподходящими", "не финальный выпуск",
                   "измеренного бизнес-прироста нет", "полнота загрузки неизвестна"):
        assert phrase in scope
    roadmap = note.tables[-1]
    assert len(roadmap.rows) == 4
    assert [cell.text for cell in roadmap.rows[0].cells] == ["Приоритет", "Следующий эксперимент", "Как оценить"]
    assert [row.cells[0].text for row in roadmap.rows[1:]] == [
        "1. Независимость", "2. Поиск и порядок", "3. Совместимость",
    ]
    bibliography = " ".join(pages[9])
    assert len(re.findall(r"(?:^|\n\n)\d+\. ", "\n\n".join(pages[9]))) == 12
    assert "полная архитектура CPC не воспроизводится" in bibliography
    assert "Статьи объясняют методы, а не подтверждают качество" in bibliography


def test_guide_preserves_structure_code_tables_and_links(tmp_path):
    source = tmp_path / "guide.md"
    source.write_text(
        "# Личный гайд\n\nВводный **текст** и `SKU`.\n"
        "## Первый шаг\n[Далее](#второй-шаг) и [файл](../config.json).\n"
        "```python\nprint('a|b')\n\nprint('next')\n```\n"
        "| Поле | Значение |\n|---|---|\n| `a|b` | x\\|y |\n"
        "### Детали\n- Проверить вход\n- Сохранить результат\n"
        "## Второй шаг\n[Источник](https://example.org/docs).\n",
        encoding="utf-8",
    )
    before = source.read_bytes()
    output = note_module.create_project_guide(source, tmp_path / "guide.docx")
    guide = Document(str(output))
    assert source.read_bytes() == before
    assert [p.text for p in guide.paragraphs if p.style.name.startswith("Heading")] == [
        "Содержание", "Первый шаг", "Детали", "Второй шаг",
    ]
    assert not guide.element.body.xpath('.//w:br[@w:type="page"]')
    assert [cell.text for cell in guide.tables[0].rows[1].cells] == ["a|b", "x|y"]
    code = [p for p in guide.paragraphs if p.style.name == "Guide Code"]
    assert [p.text for p in code] == ["print('a|b')\n\nprint('next')"]
    assert guide.styles["Guide Code"].font.name == "Courier New"
    assert any(run.font.name == "Courier New" and run.text == "SKU"
               for p in guide.paragraphs for run in p.runs)
    bookmarks = guide.element.body.xpath(".//w:bookmarkStart/@w:name")
    anchors = guide.element.body.xpath(".//w:hyperlink/@w:anchor")
    assert len(bookmarks) == len(set(bookmarks)) == 3
    assert anchors and set(anchors) <= set(bookmarks)
    urls = {r.target_ref for r in guide.part.rels.values()
            if r.reltype == RELATIONSHIP_TYPE.HYPERLINK}
    assert "https://example.org/docs" in urls
    assert (source.parent / "../config.json").resolve().as_uri() in urls
    footer = guide.sections[0].footer._element
    assert footer.xpath('.//w:fldSimple[@w:instr="PAGE"]')
    assert footer.xpath('.//w:fldSimple[@w:instr="NUMPAGES"]')


@pytest.mark.parametrize("text, message", [
    ("Без заголовка", "заголовка"),
    ("# Гайд\n```python\nx = 1", "код"),
    ("# Гайд\n| A | B |\n|---|---|\n| x |", "столбцов"),
    ("# Гайд\n[Ссылка](#нет-раздела)", "Ссылка"),
])
def test_guide_rejects_invalid_markdown_without_output(tmp_path, text, message):
    source, output = tmp_path / "guide.md", tmp_path / "guide.docx"
    source.write_text(text, encoding="utf-8")
    with pytest.raises(ValueError, match=message):
        note_module.create_project_guide(source, output)
    assert not output.exists()


def test_guide_cli_and_output_safety(tmp_path, monkeypatch, capsys):
    source, output = tmp_path / "guide.md", tmp_path / "guide.docx"
    source.write_text("# Гайд\n\n## Шаг\n\nТекст.", encoding="utf-8")
    monkeypatch.setattr(note_module, "DEFAULT_GUIDE_OUTPUT", output, raising=False)
    note_module.main(["guide", "--source", str(source)])
    assert output.exists()
    assert "Создан гайд" in capsys.readouterr().out
    for invalid in (source, tmp_path / "guide.pdf"):
        with pytest.raises(ValueError, match="отдельным"):
            note_module.create_project_guide(source, invalid)


def test_complete_personal_guide_and_worked_metrics(tmp_path):
    source = note_module.DEFAULT_GUIDE_SOURCE
    text = source.read_text(encoding="utf-8")
    blocks = note_module.parse_guide_blocks(text)
    guide = Document(str(note_module.create_project_guide(output=tmp_path / "full_guide.docx")))
    assert len(text.splitlines()) < 900
    assert len(guide.tables) == sum(kind == "table" for kind, value in blocks)
    assert [p.text for p in guide.paragraphs if p.style.name == "Guide Code"] == [
        value for kind, value in blocks if kind == "code"
    ]
    assert len(guide.element.body.xpath(".//w:bookmarkStart")) == sum(
        kind.startswith("heading_") for kind, value in blocks[1:]
    )
    urls = {r.target_ref for r in guide.part.rels.values()
            if r.reltype == RELATIONSHIP_TYPE.HYPERLINK}
    assert not any(url.startswith("file:") for url in urls)
    assert f"{note_module.REPOSITORY}/blob/main/reports/experiment_summary.json" in urls
    assert f"{note_module.REPOSITORY}/blob/main/docs/explanatory_note.md" in urls
    snapshot = json.loads((DEFAULT_SOURCE.parents[1] / "reports/experiment_summary.json").read_text())
    control = snapshot["ablation"]["test"]["e5_co_sales"]["recall@10"]
    hybrid = snapshot["test"]["models"]["learned_fusion"]["recall@10"]
    relative = f"{(hybrid / control - 1) * 100:.2f}%".replace(".", ",")
    assert relative in text and relative in DEFAULT_SOURCE.read_text(encoding="utf-8")
    ndcg_a = (1 + 1 / math.log2(5)) / (1 + 1 / math.log2(3) + 0.5)
    ndcg_b = 1 / math.log2(3)
    assert f"{ndcg_a:.4f}" in text
    assert f"{(ndcg_a + ndcg_b) / 2:.4f}" in text
    for phrase in ("вымышлены", "2 171 SKU", "121 SKU", "353 SKU", "95,04%",
                   "InfoNCE", "BPR", "PairLogit", "альтернативный ранжировщик",
                   "tools.defense presentation", "tools.defense note", "tools.defense guide",
                   "tools.defense summary", "Что измеряет эксперимент", "не предсказывают заказ",
                   "один отдельный модуль", "не универсальные границы", "общем словаре правил",
                   "сведения о брендах берутся из каталога",
                   "Модели не создают автоматически полный набор товарных ограничений"):
        assert phrase in text
    for component in ("Адаптер", "LightGCN", "CatBoostRanker"):
        status_table = text.split("Как подготовлен", 1)[1].split("### TF-IDF", 1)[0]
        row = next(line for line in status_table.splitlines() if line.startswith(f"| {component} |"))
        assert "Обучен" in row
    assert "Уже предобучена; используем готовые веса без изменения" in text
    assert "25% E5 + 25% co-sales + 25% адаптер + 25% LightGCN" in text