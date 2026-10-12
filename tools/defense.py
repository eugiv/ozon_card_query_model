"""Служебная пересборка материалов защиты; не часть ML-конвейера."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import tempfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import docx.shared
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "reports/diploma_task_definition_3slides_v5.pptx"
SUMMARY = ROOT / "reports/experiment_summary.json"
ANALYSIS = ROOT / "reports/analysis_slides_data.json"
REPOSITORY = "https://github.com/eugiv/ozon_card_query_model"
WHITE, BLACK, MINT = "FFFFFF", "142B35", "44CCA0"
PURPLE, BLUE, PINK = "7961C9", "397DBA", "DF7185"
LIGHT, GREY, DARK_MINT = "F0F5F5", "586E77", "087557"
COLORS = (GREY, BLUE, DARK_MINT, PURPLE, PINK, MINT)
FONT = "Arial"
FONT_DIR = Path("/System/Library/Fonts/Supplemental")
WIDTH, HEIGHT, SCALE = 20, 11.25, 80
TOTAL = 25
DEFAULT_SOURCE = ROOT / "docs/explanatory_note.md"
DEFAULT_OUTPUT = ROOT / "reports/explanatory_note.docx"
DEFAULT_GUIDE_SOURCE = ROOT / "docs/recommender_methodology.md"
DEFAULT_GUIDE_OUTPUT = ROOT / "reports/project_guide.docx"
PAGEBREAK = "<!-- PAGEBREAK -->"
LINK = re.compile(r"\[([^]]+)]\(((?:[^\s()]|\([^\s()]*\))+)\)")
INLINE = re.compile(r"`[^`\n]+`|" + LINK.pattern + r"|\*\*[^*\n]+\*\*|\*[^*\n]+\*")


def rgb(color):
    return RGBColor.from_string(color)


def number(value):
    return f"{int(value):,}".replace(",", " ")


def percent(value, places=2):
    return f"{value * 100:.{places}f}%".replace(".", ",")


def decimal(value, places=4):
    return f"{value:.{places}f}".replace(".", ",")


def period(start, end):
    return f"{date.fromisoformat(start):%d.%m}–{date.fromisoformat(end):%d.%m}"


def after_decision(summary, key):
    start = date.fromisoformat(summary["config"][key]) + timedelta(days=1)
    return period(start.isoformat(), (start + timedelta(days=summary["config"]["window_days"] - 1)).isoformat())


def presentation_font(size, bold=False):
    """Локальный шрифт для проверки размеров; без скачивания и привязки к macOS."""
    candidates = (
        FONT_DIR / ("Arial Bold.ttf" if bold else "Arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu") / ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2") / ("LiberationSans-Bold.ttf" if bold else "LiberationSans-Regular.ttf"),
    )
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), round(size * SCALE / 72))
    return ImageFont.truetype("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf", round(size * SCALE / 72))


def refresh_summary(config_path=ROOT / "config.json", output=SUMMARY, results_path=None):
    """Публичный снимок только агрегатов: без SKU, текстов и повторного обучения."""
    config = json.loads(config_path.read_text(encoding="utf-8"))
    folders = {key: config_path.parent / value for key, value in config["outputs"].items()}
    paths = {key: folder / "metrics.json" for key, folder in folders.items()}
    reports = {key: json.loads(path.read_text(encoding="utf-8")) for key, path in paths.items()}
    baseline, evaluation, release = (reports[key] for key in ("baseline", "evaluation", "release"))
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    metric_names = ("recall@5", "recall@10", "precision@10", "ndcg@5", "ndcg@10", "mrr")

    def cohort(models):
        selected = models[evaluation["selected_method"]]
        counts = {key: selected[key] for key in
                  ("evaluated_skus", "positive_pairs", "candidate_pairs", "retrieval_recall")}
        for model in models.values():
            if any(model[key] != selected[key] for key in
                   ("evaluated_skus", "positive_pairs", "candidate_pairs")):
                raise ValueError("Модели сравниваются на разных выборках")
        counts["models"] = {key: {metric: value[metric] for metric in metric_names}
                            for key, value in models.items()}
        return counts

    def neural(key):
        metrics = reports[key]
        fields = ("training_cutoff", "train_end", "epochs", "epochs_completed",
                  "dimensions", "layers", "sold_edges", "train_skus", "positive_pairs")
        return {field: metrics[field] for field in fields if field in metrics}

    snapshot = {
        "schema_version": 2, "run_id": "stage3_" + release["as_of"][:10].replace("-", ""),
        "updated_on": release["as_of"][:10],
        "evaluation_status": "reused_diagnostic_period",
        "source": {**evaluation["source"], "catalog_clusters": baseline["dictionary"]["clusters"],
                   "dictionary_as_of": baseline["config"]["dictionary_as_of"]},
        "aggregation": evaluation["aggregation"],
        "dictionary": {key: baseline["dictionary"][key] for key in
                       ("ambiguous_queries", "mapped_queries", "clusters", "algorithm_version_known")},
        "analysis_reference": {"rows": analysis["source"]["rows"],
                               "date_max": analysis["source"]["date_max"]},
        "config": baseline["config"],
        "training": {key: baseline["training"][key] for key in ("rows", "skus", "groups")},
        "selected_iterations": baseline["selected_iterations"],
        "selected_method": evaluation["selected_method"],
        "fusion": {key: evaluation["fusion"][key] for key in ("weights", "offset", "selection")},
        "test": cohort(evaluation["test"]["models"]),
        "zero_views_test": cohort(evaluation["test"]["zero_views"]),
        "new_days": {**cohort(evaluation["new_days"]["models"]),
                     "decision": evaluation["new_days"]["decision"],
                     "target_end": evaluation["new_days"]["target_end"],
                     "horizon_days": evaluation["new_days"]["horizon_days"]},
        "ablation": reports["ablation"],
        "neural_training": {key: neural(key) for key in
                            ("adapter", "graph", "final_adapter", "final_graph")},
        "final_training": {"rows": release["training_rows"], "skus": release["training_skus"],
                           "groups": release["training_groups"], "through": release["source_end"]},
        "current_queue": {"as_of": release["as_of"], "traffic_data_through": release["source_end"],
                          "text_snapshot_day": "2026-10-06", "catalog_as_of": release["catalog_as_of"],
                          "candidate_pairs": release["candidate_pairs"],
                          "recommended_review_pairs": release["review_pairs"],
                          "review_skus": release["review_skus"], "decisions": release["decisions"],
                          "compatibility": release["compatibility"], "top_k": release["top_k"]},
        "limitations": release["limitations"] + evaluation["limitations"],
        "provenance": {str(path.relative_to(config_path.parent)) if path.is_relative_to(config_path.parent) else path.name:
                       hashlib.sha256(path.read_bytes()).hexdigest()
                       for path in [*paths.values(), ANALYSIS, config_path]},
    }
    if snapshot["source"]["fact_rows"] != snapshot["aggregation"]["input_rows"]:
        raise ValueError("Источник и аудит расходятся")
    if results_path is not None:
        results_path = Path(results_path)
        snapshot["stage4"] = results_summary(results_path, folders["evaluation"], snapshot)
        for path in [results_path, results_path.parent / "manifest.json"]:
            key = str(path.relative_to(config_path.parent)) if path.is_relative_to(config_path.parent) else path.name
            snapshot["provenance"][key] = hashlib.sha256(path.read_bytes()).hexdigest()
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent, prefix=".summary-") as temporary:
        path = Path(temporary) / "summary.json"
        path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        path.replace(output)
    return snapshot


def results_summary(path, evaluation, snapshot):
    """Компактные проверяемые агрегаты части 4, без исходных строк и путей входов."""
    manifest = json.loads((path.parent / "manifest.json").read_text(encoding="utf-8"))
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    evaluation_digest = hashlib.sha256((evaluation / "manifest.json").read_bytes()).hexdigest()
    if manifest.get("files", {}).get(path.name) != digest or not any(
        Path(parent).resolve() == evaluation.resolve() and expected == evaluation_digest
        for parent, expected in manifest.get("parents", {}).items()
    ):
        raise ValueError("Результаты части 4 не соответствуют проверочному эксперименту")
    source = json.loads(path.read_text(encoding="utf-8"))
    if (source["status"] != "fixed_pool_diagnostic_not_independent_evaluation"
            or source["selected_method"] != snapshot["selected_method"]
            or source["fusion"] != {key: snapshot["fusion"][key] for key in ("weights", "offset")}):
        raise ValueError("Статус или модель части 4 расходятся с исходной оценкой")
    public = {key: source[key] for key in ("status", "seed", "repetitions", "limitations")}
    public["scope"] = "fixed_candidate_pool_reranking"
    buckets = {"top_5": "1-5", "rank_6_10": "6-10", "rank_11_50": "11-50",
               "rank_over_50": ">50", "missing_from_pool": "missing_from_pool"}
    for split in ("test", "new_days"):
        part = source[split]
        for method, expected in snapshot[split]["models"].items():
            actual = part["clean_comparisons"][method]
            if any(not math.isclose(actual[key], value, abs_tol=1e-12, rel_tol=0)
                   for key, value in expected.items()):
                raise ValueError("Исходные метрики части 4 расходятся с оценкой")
            if any(actual[key] != snapshot[split][key] for key in
                   ("evaluated_skus", "positive_pairs", "candidate_pairs")):
                raise ValueError("Знаменатели части 4 расходятся с оценкой")
        distribution = part["rank_distribution"]
        public[split] = {
            "horizon_days": part["horizon_days"],
            "rank_distribution": {name: {"macro_share": distribution["macro_fractions"][key],
                                         "positive_pairs": distribution["pair_counts"][key]}
                                  for name, key in buckets.items()},
            "scenarios": {},
        }
        for name, scenario in part["scenarios"].items():
            selected = scenario["models"]["selected"]
            public[split]["scenarios"][name] = {
                "recall@10": selected["recall@10"]["mean"],
                "recall@10_std": selected["recall@10"]["std"],
                "ndcg@10": selected["ndcg@10"]["mean"],
                "overlap@10": scenario["selected_top10"]["intersection_over_k"]["mean"],
                "delta_recall@10": {key: scenario["selected_delta_vs_clean"][key]
                                    for key in ("difference", "lower", "upper")},
            }
    return public


class Canvas:
    """Единая сетка и проверяемый размер текста; диаграммы остаются редактируемыми."""

    def __init__(self, presentation, title, chapter, *, cover=False):
        self.slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb(BLACK if cover else WHITE)
        self.image = Image.new("RGB", (1600, 900), "#" + (BLACK if cover else WHITE))
        self.draw = ImageDraw.Draw(self.image)
        self.text_bounds = []
        self.shape(.6, .22, .12, .35, MINT)
        self.text(chapter, .9, .25, 18.5, .35, 14, GREY if not cover else MINT, bold=True)
        if not cover:
            self.text(title, .6, .60, 18.8, 1.92, 56, bold=True)
            self.line(.6, 2.72, 19.4, 2.72, LIGHT, 2)
            self.line(.6, 2.72, 2.8, 2.72, MINT, 4)
        self.text("OZON  /  ПОИСК НОВЫХ КЛАСТЕРОВ", .6, 10.84, 16, .3, 14, WHITE if cover else GREY)
        self.text(f"{len(presentation.slides):02d} / {TOTAL}", 18, 10.84, 1.4, .3, 14, WHITE if cover else GREY, align="right")

    def text(self, value, x, y, width, height, size=22, color=BLACK, *, bold=False, align="left", url=None):
        font = presentation_font(size, bold)
        lines = []
        for original in value.split("\n"):
            current = ""
            for word in original.split(" "):
                candidate = f"{current} {word}" if current else word
                if self.draw.textlength(candidate, font=font) > width * SCALE and current:
                    lines.append(current)
                    current = word
                else:
                    current = candidate
            lines.append(current)
        line_height = size * SCALE / 72 * 1.08
        if any(self.draw.textlength(line, font=font) > width * SCALE + 1 for line in lines):
            raise ValueError(f"Текст шире поля: {value}")
        if len(lines) * line_height > height * SCALE + 1:
            raise ValueError(f"Текст выше поля ({len(lines)} строк / {height}): {value}")
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame
        frame.clear()
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(lines):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
            p.line_spacing = 1.08
            p.space_after = p.space_before = Pt(0)
            run = p.add_run()
            run.text = line
            run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
            run.font.color.rgb = rgb(color)
            if url:
                run.hyperlink.address = url
            length = self.draw.textlength(line, font=font)
            offset = {"left": 0, "right": width * SCALE - length, "center": (width * SCALE - length) / 2}[align]
            self.draw.text((x * SCALE + offset, y * SCALE + i * line_height), line, font=font, fill="#" + color, anchor="lt")
        self.text_bounds.append((x, y, width, len(lines) * line_height / SCALE, value))
        return box

    def shape(self, x, y, width, height, color=LIGHT, *, oval=False, rounded=False):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.OVAL if oval else (
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE),
                                             Inches(x), Inches(y), Inches(width), Inches(height))
        if rounded:
            shape.adjustments[0] = .12
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        etree.SubElement(shape._element.spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
        box = (x * SCALE, y * SCALE, (x + width) * SCALE, (y + height) * SCALE)
        if oval:
            self.draw.ellipse(box, fill="#" + color)
        elif rounded:
            self.draw.rounded_rectangle(box, radius=12, fill="#" + color)
        else:
            self.draw.rectangle(box, fill="#" + color)
        return shape

    def line(self, x1, y1, x2, y2, color=MINT, width=2):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.line.color.rgb, shape.line.width = rgb(color), Pt(width)
        etree.SubElement(shape._element.spPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
        self.draw.line((x1 * SCALE, y1 * SCALE, x2 * SCALE, y2 * SCALE), fill="#" + color, width=max(1, round(width)))

    def note(self, value):
        self.slide.notes_slide.notes_text_frame.text = value

    def takeaway(self, value):
        self.shape(.6, 9.80, 18.8, .78, BLACK, rounded=True)
        self.shape(.82, 9.99, .08, .40, MINT)
        self.text(value, 1.05, 9.99, 18.0, .45, 22, WHITE, bold=True)

    def label(self, value, x, y, width, color=DARK_MINT):
        self.text(value, x, y, width, .48, 24, color, bold=True)

    def bars(self, labels, values, x, y, width, height, *, colors=None, percent_axis=False, maximum=None, horizontal=False):
        data = ChartData()
        data.categories = labels
        data.add_series("Значение", values)
        chart = self.slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                            Inches(width), Inches(height), data).chart
        chart.has_legend = False
        chart.has_title = False
        plot = chart.plots[0]
        plot.gap_width = 80
        plot.has_data_labels = True
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.font.name = FONT
        plot.data_labels.font.size = Pt(18)
        plot.data_labels.font.color.rgb = rgb(BLACK)
        fmt = "0.00%" if percent_axis else "#,##0"
        plot.data_labels.number_format = fmt
        plot.data_labels.number_format_is_linked = False
        for point, color in zip(chart.series[0].points, colors or [MINT] * len(values), strict=True):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(color)
            point.format.line.fill.background()
        axis = chart.value_axis
        axis.minimum_scale = 0
        axis.maximum_scale = maximum or max(values) * 1.20 or 1
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.size = Pt(18)
        axis.tick_labels.number_format = fmt
        axis.has_major_gridlines = True
        axis.major_gridlines.format.line.color.rgb = rgb(LIGHT)
        chart.category_axis.tick_labels.font.name = FONT
        chart.category_axis.tick_labels.font.size = Pt(18)
        chart.category_axis.tick_labels.font.color.rgb = rgb(BLACK)
        if horizontal:
            chart.category_axis.reverse_order = True
            chart.category_axis.tick_labels.offset = 100
            for i, (label, value) in enumerate(zip(labels, values, strict=True)):
                row_y = y + .15 + i * (height - .4) / len(values)
                self.preview_text(label.replace("\n", " "), x, row_y, 18)
                bar_x = x + width * .32
                bar_w = value / axis.maximum_scale * width * .58
                self.shape_preview(bar_x, row_y, bar_w, .27, (colors or [MINT] * len(values))[i])
                self.preview_text(percent(value) if percent_axis else number(value), bar_x + bar_w + .1, row_y, 18)
            return chart
        # PNG — макет той же диаграммы. Нативный экспорт проверяется отдельно.
        top, base = y + .4, y + height - .8
        span = base - top
        n = len(values)
        for i, (label, value) in enumerate(zip(labels, values, strict=True)):
            col_x = x + .7 + i * (width - 1) / n
            col_w = (width - 1) / n * .55
            col_h = value / axis.maximum_scale * span
            self.shape_preview(col_x, base - col_h, col_w, col_h, (colors or [MINT] * n)[i])
            self.preview_text(percent(value, 1) if percent_axis else number(value), col_x - .1, base - col_h - .45, 18)
            self.preview_text(label, col_x - .1, base + .2, 18)
        return chart

    def shape_preview(self, x, y, width, height, color):
        self.draw.rectangle((x * SCALE, y * SCALE, (x + width) * SCALE, (y + height) * SCALE), fill="#" + color)

    def preview_text(self, value, x, y, size):
        font = presentation_font(size)
        self.draw.multiline_text((x * SCALE, y * SCALE), value, font=font, fill="#" + BLACK, spacing=4)


def apply_theme(presentation):
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for relationship in presentation.slide_master.part.rels.values():
        if relationship.reltype.endswith("/theme"):
            part = relationship.target_part
            root = etree.fromstring(part.blob)
            scheme = root.find(".//a:clrScheme", ns)
            scheme.set("name", "Тема")
            palette = {"dk1": BLACK, "lt1": WHITE, "dk2": "292B30", "lt2": LIGHT, "accent1": MINT,
                       "accent2": PINK, "accent3": PURPLE, "accent4": BLUE, "accent5": "999999", "accent6": DARK_MINT,
                       "hlink": BLUE, "folHlink": PURPLE}
            for name, color in palette.items():
                element = scheme.find("a:" + name, ns)
                for child in list(element):
                    element.remove(child)
                etree.SubElement(element, "{" + ns["a"] + "}srgbClr", val=color)
            for element in root.findall(".//a:fontScheme/a:majorFont/a:latin", ns) + root.findall(".//a:fontScheme/a:minorFont/a:latin", ns):
                element.set("typeface", FONT)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def cover(p, summary, analysis):
    c = Canvas(p, "", "ДИПЛОМНЫЙ ПРОЕКТ · 2026", cover=True)
    c.text("Рекомендательная\nсистема поисковых\nкластеров Ozon", .65, 1.40, 12.2, 4.55, 70, WHITE, bold=True)
    c.text("Находим новые поисковые запросы\nдля вашего товара.", .7, 6.25, 12.1, 1.30, 32, MINT)
    c.shape(13.30, 1.1, 6.0, 6.0, "203C48", oval=True)
    for x, y, label, color in [(13.75, 1.6, "СМЫСЛ", PURPLE), (16.5, 3.1, "ПРОДАЖИ", BLUE),
                               (13.7, 5.3, "НОВЫЕ\nЗАПРОСЫ", DARK_MINT)]:
        c.line(16.25, 4.35, x + 1.1, y + .55, MINT, 2)
        c.shape(x, y, 2.4, 1.15, color, rounded=True)
        c.text(label, x + .12, y + .25, 2.16, .75, 18, WHITE, bold=True, align="center")
    c.shape(15.1, 3.9, 2.35, 1.25, MINT, rounded=True)
    c.text("ТОВАР", 15.2, 4.30, 2.15, .45, 24, BLACK, bold=True, align="center")
    c.line(.65, 8.20, 19.35, 8.20, "3B5660", 1)
    c.text("Автор: [ФИО автора]\nПрограмма: [Название программы]", .65, 8.65, 10.8, 1.15, 22, WHITE)
    c.text(f"Руководитель: [ФИО руководителя]\nРезультаты: {date.fromisoformat(summary['updated_on']):%d.%m.%Y}", 12.1, 8.65, 7.25, 1.15, 22, WHITE)
    c.note("Рабочая гибридная система: готовая E5, обученные адаптер и LightGCN, сравнение с CatBoost. На слайдах показаны результаты восстановления продающих связей; изменение SEO остаётся бизнес-решением. Замените титульные реквизиты.")
    return c


def route(p, summary, analysis):
    c = Canvas(p, "От задачи — к проверяемому результату", "СТРУКТУРА ЗАЩИТЫ")
    items = [("01", "Постановка задачи", "Какую проблему решаем и что рекомендуем"),
             ("02", "Анализ данных", "Какие данные есть и где их ограничения"),
             ("03", "Методика решения", "Как ищем, обучаем и проверяем кластеры"),
             ("04", "Результаты", "Сравнение моделей и устойчивость к ошибкам"),
             ("05", "Выводы и заключение", "Что достигнуто и как улучшать систему"),
             ("06", "Список источников", "На чём основаны методы и расчёты")]
    for i, (n, title, description) in enumerate(items):
        x, y = .65 + (i % 2) * 9.7, 3.1 + (i // 2) * 2.15
        c.text(n, x, y, 1.2, .75, 44, DARK_MINT, bold=True)
        c.text(title, x + 1.55, y + .10, 7.5, .65, 30, bold=True)
        c.text(description, x + 1.55, y + .88, 7.4, .87, 22, GREY)
        c.line(x + 1.55, y + 1.90, x + 9.0, y + 1.90, LIGHT, 1)
    c.takeaway("От данных и обучения — к измеренному качеству и воспроизводимой выдаче.")
    c.note("Все шесть пунктов критериев: постановка, анализ, методика, результаты, выводы и заключение, список источников. Реализованный контур отделён от будущих расширений. Записка содержит десять страниц, то есть меньше предельных пятнадцати; число слайдов не является числом страниц записки.")
    return c


def problem(p, summary, analysis):
    c = Canvas(p, "Товар подходит.\nНо покупатель его не видит.", "1 · ПОСТАНОВКА ЗАДАЧИ / ПРОБЛЕМА")
    c.label("Покупатель может не знать название товара", .6, 2.95, 9.15)
    c.text("Он ищет решение своей задачи. Анализ только известных запросов не показывает все возможные связи товара с потребностями покупателя.", .6, 3.65, 8.7, 1.9, 26)
    c.label("Что такое SEO карточки?", .6, 5.80, 9)
    c.text("Работа с названием, описанием и характеристиками, чтобы товар находили по подходящим запросам.", .6, 6.5, 8.7, 1.3, 24)
    c.text("Мало просмотров — ещё не признак плохого запроса и не доказательство отсутствия фразы в карточке.", .6, 8.1, 8.7, 1.2, 20, GREY)
    c.shape(10.2, 3.05, 9.2, 1.22, LIGHT)
    c.label("меламиновая губка", 10.5, 3.27, 8.55, BLACK)
    c.text("Запрос по названию товара", 10.5, 3.82, 8.55, .4, 18, GREY)
    c.line(14.8, 4.35, 14.8, 4.80)
    c.shape(12.1, 4.8, 5.4, 1.30, MINT, oval=True)
    c.text("МЕЛАМИНОВАЯ ГУБКА", 12.3, 5.27, 5.0, .55, 22, bold=True, align="center")
    c.line(14.8, 6.15, 14.8, 6.65, PURPLE)
    c.shape(10.2, 6.65, 9.2, 1.75, LIGHT)
    c.text("губка для белой подошвы\nсредство от следов скотча", 10.5, 6.9, 8.55, 1.0, 26, PURPLE, bold=True)
    c.text("Запросы по задаче покупателя", 10.5, 7.93, 8.5, .40, 18, GREY)
    c.text("Условный пример. Совместимость с поверхностью требует проверки.", 10.2, 8.65, 9.2, .90, 18, GREY)
    c.takeaway("Ищем новые связи с потребностями, а не перестановки слов в одной фразе.")
    c.note("Принятый ранее пример сохранён. Это иллюстрация принципа, не реальный результат рекомендаций: пригодность меламиновой губки для конкретной поверхности проверяется по инструкции. Поисковые кластеры самостоятельные; варианты внутри normquery уже объединены WB. Видимость определяется не только SEO. Источник: docs/recommender_methodology.md.")
    return c


def goal(p, summary, analysis):
    c = Canvas(p, "Цель — открыть релевантный\nмалотрафиковый спрос", "1 · ПОСТАНОВКА ЗАДАЧИ / ЦЕЛЬ И ЦЕННОСТЬ")
    c.text("Для каждого товара найти отдельные кластеры normquery, которые ему подходят,\nно пока получают мало или совсем не получают просмотров его карточки.", .6, 2.97, 18.8, 1.05, 26)
    for i, (title, body) in enumerate([
        ("Опираемся на опыт", "Продающие кластеры товара и связи, обнаруженные у других товаров."),
        ("Находим новые связи", "Смысловая близость и поведение помогают отобрать и упорядочить кандидатов."),
        ("Проверяем товар", "Наличие фразы и явные конфликты бренда, объёма, материала и назначения.")]):
        x = .6 + i * 6.5
        c.text(f"0{i + 1}", x, 4.7, 1.4, .9, 48, DARK_MINT, bold=True)
        c.label(title, x, 5.80, 5.8, BLACK)
        c.text(body, x, 6.5, 5.7, 1.9, 24)
    c.text("Ценность: меньше ручного перебора, меньше дублей и более точный поисковый охват.\nИщем множество конкретных потребностей, а не общий трафик по запросу «средство для уборки».", .6, 8.4, 18.8, 1.05, 22, GREY)
    c.takeaway("Текущий результат — список на проверку. Карточки автоматически не меняем.")
    c.note("Кластер normquery — готовая каноническая группа запросов WB. Текущий текст используется отдельно от исторического обучения. Консервативные правила материала и назначения реализованы, но охватывают только явно названные свойства из ограниченного словаря. Отсутствие конфликта не доказывает полную совместимость. Результат — гипотезы review. Суммы частотностей не равны числу новых покупателей.")
    return c


def metrics(p, summary, analysis):
    c = Canvas(p, "Как измеряем качество рекомендаций", "1 · ПОСТАНОВКА ЗАДАЧИ / КРИТЕРИИ")
    c.text("Список строим по прошлому. Затем проверяем заказы в следующем периоде.\nK — число первых рекомендаций для оценки, а не обязательная длина рабочего списка.", .6, 2.92, 18.8, .95, 24)
    c.label("Условный пример: меламиновая губка", .6, 4.10, 9)
    c.text("В будущем заказали по 3 кластерам:\nбелая подошва, чистка кафеля, пол.\nПервые два найдены на местах 1 и 3;\nкластер «для пола» не найден.", .6, 4.78, 8.4, 2.3, 24)
    c.text("Первые 5 рекомендаций", 10.2, 4.10, 9.2, .5, 24, PURPLE, bold=True)
    c.text("1 · губка для белой подошвы — заказ\n2 · средство от следов скотча — нет заказа\n3 · жесткая губка для чистки кафеля — заказ\n4 · губка для чистки ванны — нет заказа\n5 · губка для чистки мебели — нет заказа", 10.2, 4.78, 9.2, 2.05, 20)
    for x, title, formula, example in [
        (.6, "Recall · полнота", "Recall@K = H / P", "2 из 3 → 66,7%"),
        (7.05, "Precision · подтверждение", "Precision@K = H / K", "2 из 5 → 40%"),
        (13.5, "NDCG · порядок", "NDCG@K = DCG / IDCG", "1,5 / 2,131 → 70,4%")]:
        c.line(x, 7.20, x + 5.9, 7.20, MINT, 3)
        c.label(title, x, 7.47, 5.9, BLACK)
        c.text(formula, x, 8.13, 5.9, .5, 22)
        c.text(example, x, 8.78, 5.9, .55, 26, DARK_MINT, bold=True)
    c.text("DCG = Σ (i=1…K) r(i)/log₂(i+1);  IDCG = Σ (i=1…min(K,P)) 1/log₂(i+1)", .6, 9.60, 18.8, .45, 20)
    c.text("P — все кластеры с заказом; H — найденные; r(i) = 1 при заказе, иначе 0. Нет заказа ≠ не подходит.", .6, 10.17, 18.8, .4, 18, GREY)
    c.note("Пример полностью условный, названия сохранены из принятой версии. Формулы рассчитаны для одного SKU: DCG=1+1/log2(4)=1.5; IDCG=1+1/log2(3)+1/log2(4)=2.13093; NDCG=0.70392. В коде вычисляем K=5/10, усредняем по всем SKU с положительными целями, включая SKU без найденных кандидатов. Пропущенные цели входят в P. Наблюдаемая Precision не равна экспертной релевантности. Диагностика дополнительно использует полноту общего пула и MRR. Код: recommendation_metrics.py.")
    return c


def data_slide(p, summary, analysis):
    c = Canvas(p, "Данные: вся доступная история Ozon", "2 · АНАЛИЗ / ИСТОЧНИК И ОСОБЕННОСТИ")
    source = summary["source"]
    for x, value, label in [(.6, number(source["fact_rows"]), "дневных записей"),
                            (7.15, number(source["skus"]), "разных товаров"),
                            (13.7, period(source["first_day"], source["last_day"]), "период истории · 2026")]:
        c.text(value, x, 2.95, 5.7, .95, 48, DARK_MINT, bold=True)
        c.text(label, x, 4.00, 5.7, .48, 22, GREY)
    c.label("Сколько разных товаров наблюдалось в каждом месяце", .6, 4.8, 18.8, BLACK)
    monthly = analysis["source"]["monthly"][:9]
    c.bars(["янв*", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен"],
           [row["sku_count"] for row in monthly], .6, 5.4, 12, 3.75, colors=[MINT] * 7 + [PURPLE, PURPLE])
    c.text("Одна строка: день + товар + запрос.\nПоиски — запрос искали.\nПросмотры — карточку увидели,\nэто не клики.\nЗаказы — количество покупок.", 13.2, 5.47, 6.2, 2.8, 22)
    c.text("* Январь начинается 9-го числа.\nОктябрь неполный, поэтому не показан.", 13.2, 8.40, 6.2, .9, 18, GREY)
    c.takeaway("В августе и сентябре меньше наблюдаемых SKU. Это не график падения продаж.")
    c.note("Источник default_staging.stg_ozon_product_queries_details. Месячные значения — число SKU с любыми фактами, не объём спроса, продаж или полный каталог. Причина изменения охвата неизвестна. Наличие дня не доказывает полноту загрузки. Суммы unique_search_users и unique_view_users по дням/вариантам не являются уникальными людьми за весь период. Частотность дедуплицируется по query/date через max. Источник графика: reports/analysis_slides_data.json, source.monthly.")
    return c


def quality(p, summary, analysis):
    c = Canvas(p, "Словарь сохраняет большую часть заказов", "2 · АНАЛИЗ / ПОКРЫТИЕ И АНОМАЛИИ")
    audit = summary["aggregation"]
    c.text("Исходные фразы сопоставляем с normquery — готовыми кластерами WB.\nОдна связь — товар и исходная фраза за всю историю, без повторов по дням.", .6, 2.98, 18.8, .95, 24)
    rows = [("Связи «товар + исходная фраза»", audit["mapped_pairs"], audit["raw_pairs"], BLUE),
            ("Заказы по всей истории", audit["covered_orders"], audit["orders"], MINT)]
    for i, (label, covered, total, color) in enumerate(rows):
        y = 4.38 + i * 2.08
        share = covered / total
        c.label(label, .6, y, 14, BLACK)
        c.text(percent(share), 16.3, y, 3.1, .65, 32, DARK_MINT, bold=True, align="right")
        c.shape(.6, y + .7, 18.8, .40, LIGHT)
        c.shape(.6, y + .7, 18.8 * share, .40, color)
        c.text(f"{number(covered)} из {number(total)} сопоставлены", .6, y + 1.32, 12, .5, 22)
        c.text(f"{number(total - covered)} — без кластера", 13.0, y + 1.32, 6.4, .5, 20, GREY, align="right")
    missing = summary["source"]["missing_days"]
    days = ", ".join(f"{date.fromisoformat(day):%d.%m}" for day in missing)
    c.text(f"{number(summary['dictionary']['ambiguous_queries'])} фраз имели несколько назначений: выбираем одно на дату снимка.\nДней без записей: {len(missing)} · {days}. Пропуски не подменяем нулями.", .6, 8.72, 18.8, .85, 20, GREY)
    c.takeaway(f"Сопоставление сохраняет {percent(audit['order_coverage'])} заказов для обучения и проверки.")
    c.note(f"Актуальный аудит по 07.10: {audit['mapped_pairs']}/{audit['raw_pairs']} сопоставленных пар; {audit['covered_orders']}/{audit['orders']} заказов, разные знаменатели. {audit['orders_without_views_daily_pairs']} дневных пар с заказом и нулём просмотров — диагностический флаг, не автоматическое удаление. loaded_at не является версией алгоритма. Источник: experiment_summary.json.aggregation.")
    return c


def dependencies(p, summary, analysis):
    c = Canvas(p, "Малый трафик — зона неизвестности", "2 · АНАЛИЗ / ЗАВИСИМОСТИ")
    obs = analysis["observation"]
    c.text(f"05–18 сентября 2026 · {obs['skus']} товаров · {number(obs['pairs'])} связей «товар + кластер».\nСравниваем видимость и наличие хотя бы одного заказа в этом же периоде.", .6, 2.97, 18.8, .95, 24)
    buckets = obs["view_buckets"][1:]
    c.label("Число связей по суммарным просмотрам", .6, 4.38, 9.1, BLACK)
    c.label("Доля связей с хотя бы одним заказом", 10.3, 4.38, 9.1, BLACK)
    labels = [row["label"] for row in buckets]
    c.bars(labels, [row["pairs"] for row in buckets], .6, 5.15, 9.1, 3.3, colors=[MINT, BLUE, PURPLE, PINK])
    c.bars(labels, [row["order_pair_share"] for row in buckets], 10.3, 5.15, 9.1, 3.3,
           colors=[MINT, BLUE, PURPLE, PINK], percent_axis=True, maximum=1)
    c.text("Просмотров за 14 дней", .6, 8.64, 9.1, .4, 18, GREY)
    c.text("Это доля связей, не конверсия пользователей", 10.3, 8.64, 9.1, .4, 18, GREY)
    low = buckets[0]
    c.text(f"{low['label']} просмотров: {number(low['with_orders'])} из {number(low['pairs'])} связей получили заказ = {percent(low['order_pair_share'], 1)}.\nНулевые просмотры при заказах учитываем отдельно: источник не содержит кликов.", .6, 9.20, 18.8, .75, 20)
    c.note("Источник: observation.view_buckets в analysis_slides_data.json. Это описательная связь, не причинный эффект повышения видимости. В группах показаны суммарные просмотры, а не уникальные посетители за период. Низкая доля наблюдённых заказов не даёт достоверной отрицательной метки. Отдельно есть 17 SKU×cluster с views=0 и orders>0; их не используем как обычную воронку кликов. Значения:9177/346,2663/316,1566/471,844/664.")
    return c


def requirements_slide(p, summary, analysis):
    c = Canvas(p, "Что анализ требует от решения", "2 · АНАЛИЗ / ВЫВОДЫ ДЛЯ МЕТОДИКИ")
    rows = [("Есть много малонаблюдаемых связей", "Ищем новые кластеры, а не только оцениваем известные."),
            ("Нет заказа ≠ неверный запрос", "Неизвестные пары используем только как слабые сравнения."),
            ("Словарь датирован 26 июня", "Решения проверяем после снимка; ранняя история — прошлый опыт."),
            ("Тексты карточек — с 6 октября", "Проверяем текущую выдачу; не подмешиваем их в сентябрьское обучение.")]
    for i, (fact, action) in enumerate(rows):
        y = 3.0 + i * 1.55
        c.text(fact, .6, y, 8.8, .85, 28, bold=True)
        c.text("→", 9.5, y, .7, .5, 28, DARK_MINT)
        c.text(action, 10.35, y, 9.05, 1.10, 24)
        c.line(.6, y + 1.30, 19.4, y + 1.30, LIGHT, 1)
    c.takeaway("Текстово-поведенческий контур работает без фото и ручной разметки.")
    c.note("Анализ задаёт конкретные требования: safe many-to-one mapping; контроль пропусков; asof-словарь; temporal validation; неизвестные не являются надёжными negatives; product checks отделены от исторического pipeline. Структура источника ограничивает холодный старт. Фото/OCR перенесены после доводки товарной релевантности. Источник: docs/recommender_methodology.md, разделы о текущем тексте и границах отсева.")
    return c


def pipeline(p, summary, analysis):
    c = Canvas(p, "Два потока: оценка и текущая выдача", "3 · МЕТОДИКА РЕШЕНИЯ / АРХИТЕКТУРА")
    c.label("Исторический эксперимент", .6, 3.0, 18.8)
    for i, (title, detail) in enumerate([("Подготовка", "Факты + словарь\nSKU × normquery"), ("Поиск", "Текст + смысл\n+ совместные продажи"),
                                        ("Обучение", "Адаптер, граф\nи CatBoost"), ("Оценка", "Будущий период\nи вклад каналов")]):
        x = .6 + i * 4.95
        c.shape(x, 3.85, 3.95, 1.9, LIGHT, rounded=True)
        c.shape(x, 3.4, .66, .66, [BLUE, PURPLE, DARK_MINT, PINK][i], oval=True)
        c.text(str(i + 1), x, 3.52, .66, .4, 19, WHITE, bold=True, align="center")
        c.label(title, x + .2, 4.07, 3.55, BLACK)
        c.text(detail, x + .2, 4.75, 3.55, .80, 20)
        if i < 3:
            c.text("→", x + 4.14, 4.47, .65, .5, 26, DARK_MINT)
    c.label("Текущая выдача", .6, 6.45, 18.8)
    c.text("Финальное обучение\nпо всей истории", .6, 7.25, 5.2, 1.15, 28, bold=True)
    c.text("→", 6.0, 7.3, .65, .5, 28, DARK_MINT)
    c.text("Текст и свойства\nкарточки на дату", 7.1, 7.23, 5.5, 1.15, 28, bold=True)
    c.text("→", 13.0, 7.3, .65, .5, 28, DARK_MINT)
    c.text("Список предложений\nдля SEO-специалиста", 14.1, 7.23, 5.3, 1.15, 28, bold=True)
    c.text(f"Собственные прошлые продажи исключены. Для кандидата — не более {summary['config']['max_views']} просмотров за {summary['config']['window_days']} дней.", .6, 8.90, 18.8, .8, 22, GREY)
    c.takeaway("Низкий трафик — условие поиска. Наличие фразы и совместимость проверяем отдельно.")
    c.note("0–5 суммарных просмотров за 14 дней — экспериментальные параметры. Все прошлые продающие кластеры исключены из кандидатов. Проверочные нейросети обучены до 04.09. После оценки выполнен отдельный refit по 07.10. Эти финальные веса нельзя оценивать на сентябрьских ответах. LightGCN и адаптер обучены; cross-encoder, фото и OCR не используются. Холодный старт не решён. Код: cluster_data.py, cluster_recommender.py, full_history_recommender.py, finalize_recommender.py, seo_recommendation_policy.py.")
    return c


def search(p, summary, analysis):
    c = Canvas(p, "Смысл и поведение: четыре канала гибрида", "3 · МЕТОДИКА РЕШЕНИЯ / ПОИСК")
    items = [("Предобученная E5", "Готовая нейросеть смысла; используем её обученные веса.", BLUE),
             ("Обученный адаптер", "Уточняет векторы E5 по продающим связям нашего каталога.", PURPLE),
             ("Совместные продажи", "Другие товары продаются по связанным кластерам.", DARK_MINT),
             ("Обученный LightGCN", "Передаёт информацию по графу товаров и запросов.", PINK)]
    for i, (title, detail, color) in enumerate(items):
        y = 3.0 + i * 1.42
        c.shape(.6, y, .08, 1.1, color)
        c.text(title, .9, y, 6.4, .60, 30, color, bold=True)
        c.text(detail, 7.6, y + .08, 9.2, .85, 24)
        c.line(17.15, y + .5, 18.7, 6.3, MINT)
    c.text("ОБЩИЙ\nПУЛ", 17.4, 6.53, 2.0, 1.0, 26, DARK_MINT, bold=True, align="center")
    c.text(f"До {summary['config']['per_source']} допустимых кандидатов от канала; дубли убираем. TF-IDF и популярность — контроль.\nВсе методы затем сравниваем на одинаковом расширенном пуле.", .6, 9.02, 18.8, .75, 22, GREY)
    c.takeaway("Обучение адаптирует семантические и графовые связи к данным нашего каталога.")
    c.note("E5 — предобученная intfloat/multilingual-e5-small. Готовые веса фиксированы: замороженные веса с закреплённой версией используются без изменения. На данных проекта обучен отдельный residual MLP-адаптер поверх векторов E5; это обучение адаптера, а не fine-tuning Transformer. Co-sales рассчитывает связи непосредственно из совместных продаж. LightGCN — отдельно обученная графовая нейросеть. Её поисковый профиль строится из seed-кластеров, а не собственного SKU embedding. TF-IDF и популярность остаются контрольными каналами; извлечение и сравнение используют общий пул. Фото, OCR и cross-encoder не входят в текущую реализацию. Отсутствие сигнала графа не означает плохой запрос.")
    return c


def learning(p, summary, analysis):
    c = Canvas(p, "Учимся ставить перспективные связи выше", "3 · МЕТОДИКА РЕШЕНИЯ / ОБУЧЕНИЕ")
    c.label("1 · Будущая новая продажа", .6, 3.05, 9.1)
    c.text("В прошлом заказов по связи нет.\nВ следующем периоде появился заказ.\nЭто наблюдаемый положительный пример.", .6, 3.85, 9.1, 1.8, 26)
    c.label("2 · Восстановление скрытой связи", 10.3, 3.05, 9.1, PURPLE)
    c.text("Скрываем один продающий кластер.\nУдаляем его из истории товара и графа.\nИщем по оставшимся ориентирам.", 10.3, 3.85, 9.1, 1.8, 26)
    c.line(.6, 5.85, 19.4, 5.85, LIGHT)
    c.text("CatBoostRanker · PairLogit", .6, 6.20, 9.1, .6, 32, bold=True)
    c.text("Положительный пример выше неизвестного сравнения. Сходство, граф и общий спрос — признаки; собственный трафик не задаёт приоритет.", .6, 7.10, 9.1, 1.65, 24)
    c.text("Гибрид обратных рангов · RRF", 10.3, 6.20, 9.1, .6, 32, bold=True)
    c.text(f"score = Σ w / ({summary['fusion']['offset']} + позиция)\nВеса и метод выбираем на отдельном периоде. Это порядок, не вероятность.", 10.3, 7.10, 9.1, 1.65, 24)
    training = summary["training"]
    c.takeaway(f"Обучение: {number(training['rows'])} строк · {training['skus']} SKU · {training['groups']} групп «SKU–задача».")
    c.note("Неизвестные пары — слабые сравнения, не подтверждённые отрицательные. До 12 сложных и случайных сравнений на positive, seed=7, нормировка веса SKU–задачи. Пропущенная цель не добавляется в пул. Выбрано 600 из 100/300/600 итераций по validation NDCG@10. В RRF смещение 60, шаг весов 0,25; выбраны равные веса E5/co-sales/адаптера/LightGCN. CatBoost сравнивается отдельно. Маскировка — задача обучения ранжировщика: собственное ребро удаляется из истории и исходного co-sales графа. Это не независимая проверка нейросетей, которые могли видеть скрытые связи. Главная оценка обученных каналов — временная проверка. Источник: dormant_learning.py, experiment_summary.json.training.")
    return c


def trained_models(p, summary, analysis):
    c = Canvas(p, "Три модели обучены на данных проекта", "3 · МЕТОДИКА РЕШЕНИЯ / МОДЕЛИ И ПАРАМЕТРЫ")
    c.text("Сначала обучаем проверочные модели на прошлом. После измерения качества\nпереобучаем финальные модели на всей доступной истории по 07.10.2026.", .6, 2.97, 18.8, 1.0, 24)
    final = summary["final_training"]
    adapter = summary["neural_training"]["final_adapter"]
    graph = summary["neural_training"]["final_graph"]
    items = [
        ("CatBoostRanker", BLUE, "Сравнивает кандидатов\nи учится их упорядочивать.",
         f"PairLogit · {summary['selected_iterations']} итераций", f"{number(final['rows'])} строк\n{final['skus']} товаров"),
        ("Адаптер E5", PURPLE, "Уточняет готовые векторы\nсмысла по нашим продажам.",
         "384 → 128 → 384 · InfoNCE", f"{number(adapter['positive_pairs'])} пар\n{adapter['epochs']} эпох обучения"),
        ("LightGCN", DARK_MINT, "Учит связи в графе\nтоваров и запросов.",
         f"{graph['dimensions']} измерения · {graph['layers']} слоя · BPR", f"{number(graph['sold_edges'])} связей\n{graph['epochs_completed']} эпох обучения"),
    ]
    for i, (name, color, role, architecture, counts) in enumerate(items):
        x = .6 + i * 6.5
        c.line(x, 4.37, x + 5.8, 4.37, color, 4)
        c.text(name, x, 4.72, 5.8, .60, 30, color, bold=True)
        c.text(role, x, 5.62, 5.8, 1.35, 24)
        c.text(architecture, x, 7.2, 5.8, .85, 22, GREY)
        c.text(counts, x, 8.25, 5.8, 1.10, 26, color, bold=True)
    c.takeaway("Готовая E5 + обучение на своём каталоге: адаптер, LightGCN и CatBoostRanker.")
    c.note("Реальные финальные размеры из experiment_summary.json. Проверочный CatBoost: 37 245 строк / 334 SKU / 2 504 группы. Финальный: 46 787 / 413 / 3 053. Проверочный адаптер: 20 740 пар / 549 SKU; финальный: 23 569 / 589. Архитектура residual MLP normalize(x+0.1*MLP(x)), contrastive InfoNCE. LightGCN: 9 525 проверочных и 10 742 финальных продающих рёбер, 32 измерения, 2 слоя, BPR и L2. Проверочные нейросети обучены по 04.09, финальные по 07.10 после оценки. Падение training loss показывает обучение, но не заменяет Recall/NDCG. LightGCN не использует собственный SKU embedding при выдаче; профиль составляется из известных seed-кластеров. E5 — готовая предобученная нейросеть с фиксированными весами; на нашем каталоге обучены адаптер, LightGCN и CatBoostRanker. Итоговый гибрид выбран по validation: E5, co-sales, адаптер и LightGCN по 25%. CatBoost сохранён как обученный альтернативный ранжировщик. Cross-encoder относится к дальнейшим экспериментам, не к действующей архитектуре.")
    return c


def temporal(p, summary, analysis):
    c = Canvas(p, "Прошлое для обучения.\nСледующий период — для проверки.", "3 · МЕТОДИКА РЕШЕНИЯ / РАЗДЕЛЕНИЕ ДАННЫХ")
    c.text(f"История известна только до даты решения. Ответ — заказ в следующие {summary['config']['window_days']} дней.\nФиксируем словарь на {date.fromisoformat(summary['config']['dictionary_as_of']):%d.%m.%Y}; тексты карточек используются при выдаче.", .6, 2.99, 18.8, 1.10, 24)
    rows = [("Обучение", "История до 04.09 · пять дат решений", "CatBoost, адаптер и LightGCN; окна ответов по 14 дней", BLUE),
            ("Подбор метода", after_decision(summary, "validation_decision"), "Выбираем итерации и веса гибрида", PURPLE),
            ("Проверка", after_decision(summary, "test_decision"), "Измеряем результат выбранного метода", DARK_MINT)]
    for i, (label, dates, explanation, color) in enumerate(rows):
        y = 4.7 + i * 1.43
        c.shape(.6, y, .10, 1.05, color)
        c.text(label, .95, y + .1, 4.2, .65, 30, color, bold=True)
        c.text(dates, 5.4, y, 14.0, .55, 26, bold=True)
        c.text(explanation, 5.4, y + .64, 14.0, .50, 22, GREY)
    c.text("Сентябрь — историческая диагностика; октябрь — дополнительная короткая проверка.\nПосле измерения качества — отдельное финальное обучение для выпуска рекомендаций.", .6, 9.13, 18.8, .80, 22)
    c.note("Даты решений CatBoost: 26.06, 03.07, 07.08, 14.08, 21.08. Шаг 7 дней, ответы по 14 дней; перекрывающиеся targets не независимы. Решения 10/17/24/31.07 пропущены из-за отсутствия 21.07. Две задачи маскировки на решение. Адаптер и граф для проверки используют только историю по 04.09. Validation: 05–18.09; test: 19.09–02.10. Дополнительный период 04–07.10, параметры не выбирались по его ответам. Полный refit по 07.10 после оценки включает прежние ответы и не является моделью для сентябрьского test. Ранняя история используется как прошлый контекст, но словарь от 26.06 нельзя незаметно применять к решениям января. Все пропущенные цели остаются в знаменателе. Это не независимый test, поскольку сентябрьский период просмотрен ранее.")
    return c


def results(p, summary, analysis):
    c = Canvas(p, "Гибрид находит больше будущих\nпродающих кластеров в первой десятке", "4 · РЕЗУЛЬТАТЫ / СРАВНЕНИЕ МЕТОДОВ")
    test = summary["test"]
    c.text(f"Одна малотрафиковая проверка: {test['evaluated_skus']} SKU · {test['positive_pairs']} будущих положительных связей.\nRecall@10 — доля таких кластеров, найденных среди первых 10, в среднем по SKU.", .6, 2.99, 18.8, .95, 24)
    ids = ["lexical", "co_sales", "pretrained_embeddings", "semantic_adapter", "lightgcn", "catboost_pu", "learned_fusion"]
    labels = ["TF-IDF", "Совместные\nпродажи", "E5", "Адаптер", "LightGCN", "CatBoost", "Гибрид"]
    values = [test["models"][key]["recall@10"] for key in ids]
    c.bars(labels, values, .6, 4.38, 12.3, 4.80, colors=[GREY, BLUE, PURPLE, PINK, DARK_MINT, BLUE, MINT], percent_axis=True, maximum=.30, horizontal=True)
    best = test["models"][summary["selected_method"]]
    c.text(percent(best["recall@10"]), 13.55, 4.83, 5.85, 1.3, 64, DARK_MINT, bold=True)
    c.text(f"NDCG@10 = {decimal(best['ndcg@10'])}\nPrecision@10 = {percent(best['precision@10'])}", 13.55, 6.45, 5.85, 1.2, 26)
    c.text("Находит запросы, по которым\nпозднее появились заказы.", 13.55, 8.04, 5.85, .9, 22, GREY)
    c.takeaway("Гибрид — лучший из сравниваемых методов на данной исторической проверке.")
    c.note(f"Источник: experiment_summary.json.test, реальные средние по {test['evaluated_skus']} SKU и {test['positive_pairs']} будущим положительным связям; одинаковый пул {number(test['candidate_pairs'])}. Отсутствующие кандидаты остаются в знаменателе. Гибрид выбран до test по validation NDCG@10. Precision@10={percent(best['precision@10'])} — подтверждение будущими заказами, не экспертная релевантность. Сентябрьский период уже использован; это диагностика. Популярность — контроль, Recall@10={percent(test['models']['popularity']['recall@10'])}. Все модели реально обучены или построены; улучшение не означает прирост продаж.")
    return c


def neural_value(p, summary, analysis):
    c = Canvas(p, "Что дают обученные компоненты", "4 · РЕЗУЛЬТАТЫ / ВКЛАД МОДЕЛЕЙ")
    a = summary["ablation"]
    c.text("Одинаковые товары, ответы и пул кандидатов. Меняем только каналы ранжирования.\nВеса каждого варианта заранее выбраны на периоде подбора модели.", .6, 2.99, 18.8, .95, 24)
    c.label("Первая десятка · Recall@10", .6, 4.4, 11.9, BLACK)
    keys = ["e5_co_sales", "plus_adapter", "plus_lightgcn", "all_four"]
    c.bars(["E5 + совместные\nпродажи", "+ адаптер", "+ LightGCN", "Все четыре"],
           [a["test"][key]["recall@10"] for key in keys], .6, 5.05, 11.9, 3.55,
           colors=[GREY, PURPLE, BLUE, MINT], percent_axis=True, maximum=.30)
    diff = a["test_full_minus_control"]
    c.text("+" + percent(diff["difference"]).removesuffix("%") + " п.п.", 13.4, 4.85, 6, 1.1, 52, DARK_MINT, bold=True)
    c.text(f"95%-й интервал разницы:\n+{percent(diff['lower']).removesuffix('%')}…+{percent(diff['upper']).removesuffix('%')} п.п.", 13.4, 6.28, 6, 1.15, 24)
    control = a["test"]["e5_co_sales"]["recall@10"]
    learned = a["test"]["all_four"]["recall@10"]
    c.text(f"{percent((learned - control) / control)} относительно\nконтрольного гибрида.\nОбучение улучшает отбор.", 13.4, 7.79, 6, 1.35, 22, GREY)
    zero = summary["zero_views_test"]
    c.takeaway(f"Без прошлых просмотров: Recall@10 = {percent(zero['models'][summary['selected_method']]['recall@10'])} · {zero['evaluated_skus']} SKU / {zero['positive_pairs']} связей.")
    c.note("Источник: experiment_summary.json.ablation. Контроль E5+co-sales и все варианты используют один расширенный пул из 166 291 кандидата, 121 SKU/260 целей. Веса подбираются только на validation. Парный SKU-bootstrap, 2 000 повторов, seed7. Интервал поддерживает диагностический прирост, но не превращает повторно использованный период в независимый. Отдельно zero_views_test94SKU180целей, без постфактум фильтрации выдачи и без вывода об отсутствии в SEO. Старые 18,23% на другом пуле не являются корректным прямым контролем этого эксперимента.")
    return c


def new_period(p, summary, analysis):
    c = Canvas(p, "Новые дни: проверка без подбора настроек", "4 · РЕЗУЛЬТАТЫ / ОТДЕЛЬНЫЙ ПЕРИОД")
    new = summary["new_days"]
    c.text(f"04–07 октября 2026 · {new['horizon_days']} дня · {new['evaluated_skus']} товаров · {new['positive_pairs']} новых продающих связей.\nИспользуем уже выбранные настройки. По октябрьским ответам модель не подбирали.", .6, 2.99, 18.8, .95, 24)
    ablation = summary["ablation"]
    values = [new["models"]["lexical"]["recall@10"], ablation["new_days"]["e5_co_sales"]["recall@10"],
              new["models"][summary["selected_method"]]["recall@10"]]
    c.bars(["TF-IDF", "E5 + совместные\nпродажи", "Обученный\nгибрид"], values,
           .6, 4.50, 10.3, 4.10, colors=[GREY, BLUE, MINT], percent_axis=True, maximum=.25)
    c.label("Перенос на более поздние данные", 11.55, 4.55, 7.85, BLACK)
    c.text(f"Recall@10 гибрида — {percent(values[-1])}.\nРекомендации проверены\nбез изменения выбранной модели.", 11.55, 5.36, 7.85, 1.40, 26)
    diff = ablation["new_days_full_minus_control"]
    c.text(f"Разница с контролем: +{percent(diff['difference']).removesuffix('%')} п.п.\n95%-й интервал: {percent(diff['lower']).removesuffix('%')}…+{percent(diff['upper']).removesuffix('%')} п.п.\nВключает ноль: это короткая проверка,\nне вывод об устойчивом преимуществе.", 11.55, 7.05, 7.85, 2.10, 22, GREY)
    c.takeaway("Дополнительный опыт: выбранная система находит новые продающие связи и в октябре.")
    c.note("Источники: experiment_summary.json.new_days и ablation.new_days. Решение03.10, ответы04–07.10, 96SKU132цели, одинаковый пул215719. Параметры и метод заранее выбраны по сентябрьскому validation, а не по этим ответам. Контроль15,14%, гибрид17,95%, разница+2,81п.п., парныйSKUbootstrap95%-2,08…+7,92п.п. Интервал включает0. Нельзя приравнивать этот4-дневный горизонт к сентябрьскому14-дневному. Для устойчивости нужна новая полноценная проверка; прирост продаж и экспертная релевантность не измерены.")
    return c


def robustness(p, summary, analysis):
    c = Canvas(p, "Что происходит, если сигналы ошибаются", "4 · РЕЗУЛЬТАТЫ / ШУМ И ВЫБРОСЫ")
    stage = summary["stage4"]
    part = stage["test"]
    scenarios = part["scenarios"]
    keys = ["clean", "score_noise_05", "score_noise_15", "counter_outliers", "candidate_loss_10", "candidate_loss_30"]
    values = [scenarios[key]["recall@10"] for key in keys]
    c.text(f"Те же {summary['test']['evaluated_skus']} товаров и {summary['test']['positive_pairs']} правильных ответов. Меняем входные сигналы, а не ответы.\nМодели и веса заморожены; случайные искажения повторены {stage['repetitions']} раза.", .6, 2.98, 18.8, .98, 24)
    c.label("Доля ответов в первой десятке · Recall@10", .6, 4.05, 12.1, BLACK)
    c.bars(["Исходная\nвыдача", "Шум\n5%", "Шум\n15%", "Выбросы\n×100", "Потеря\n10% пула", "Потеря\n30% пула"],
           values, .6, 4.50, 12.1, 4.25, colors=[MINT, BLUE, PURPLE, PINK, GREY, DARK_MINT],
           percent_axis=True, maximum=max(.30, max(values) * 1.2))
    c.label("Что именно искажаем", 13.35, 4.52, 6.05, BLACK)
    c.text("Шум — ошибки оценок близости:\n5% и 15% их разброса у товара.\n\nВыброс — счётчики ×100\nу 1% случайных кандидатов.\n\nПотеря — случайно удалённая\nчасть найденных запросов.", 13.35, 5.30, 6.05, 3.67, 22)
    c.takeaway("Это проверка ранжирования найденного пула, не всей системы поиска и обучения.")
    c.note("Источник: experiment_summary.json.stage4.test.scenarios. Столбцы — macro Recall@10 гибрида, среднее трёх повторов seed7/8/9, не процент релевантных фраз. Шум 5/15% от стандартного отклонения оценок внутри SKU; поведенческий score — мультипликативный шум. Выбросы x100 глобальных продаж, спроса и поведенческого score у случайного1%; log-счётчики разворачиваются перед умножением. Собственные просмотры и условие допуска неизменны. При случайной потере10/30% цели не удаляются из знаменателя. Никакого обучения/подбора настроек. Это не опечатки в тексте и не переобучение LightGCN на шумных исходных заказах. Сентябрьский период диагностический; интервалы и совпадение первой десятки доступны в агрегатах отчёта.")
    return c


def ranking_diagnosis(p, summary, analysis):
    c = Canvas(p, "Где теряются правильные поисковые кластеры", "4 · РЕЗУЛЬТАТЫ / РАЗБОР КАЧЕСТВА")
    distribution = summary["stage4"]["test"]["rank_distribution"]
    items = [("top_5", "В первых пяти", MINT), ("rank_6_10", "На местах 6–10", DARK_MINT),
             ("rank_11_50", "На местах 11–50", BLUE), ("rank_over_50", "Ниже 50-го места", PURPLE),
             ("missing_from_pool", "Не найдены поиском", PINK)]
    c.text(f"Берём все кластеры, по которым позже появились заказы, и смотрим их позиции.\nДоли сначала считаем внутри товара, затем усредняем по тем же {summary['test']['evaluated_skus']} SKU.", .6, 2.99, 18.8, .97, 24)
    for i, (key, title, color) in enumerate(items):
        y = 4.43 + i * .93
        value = distribution[key]["macro_share"]
        c.text(title, .6, y, 6.1, .5, 24)
        c.shape(7.05, y + .10, 7.0, .25, LIGHT)
        if value > 0:
            c.shape(7.05, y + .10, 7.0 * value, .25, color)
        c.text(percent(value), 14.45, y - .03, 2.3, .56, 27, color, bold=True, align="right")
        c.text(f"{distribution[key]['positive_pairs']} связей", 17.12, y + .02, 2.28, .52, 19, GREY, align="right")
    c.text("Проценты — среднее по товарам; числа справа — все связи выборки.\nЭто разные способы усреднения: число связей / 260 не равно показанному проценту.", .6, 8.95, 18.8, .70, 19, GREY)
    c.takeaway("Не найденный кластер не исправить сортировкой. Найденный ниже десятки — задача ранжирования.")
    c.note("Источник: experiment_summary.json.stage4.test.rank_distribution. Полный набор121SKU/260positivepairs; порядок learned_fusion score descending, normquery ascending. Пять взаимоисключающих корзин; macro-доли суммируются в1. Первые две суммарно равны исходномуRecall@10, последняя равна1-retrieval_recall. Отдельный целевой query может не входить в общий пул — он не удаляется из оценки. Это распределение наблюдаемых будущих positive, не полный каталог релевантных фраз. Неизвестные без заказа не объявляются негативами. Практический вывод касается диагностики поиска/ранжирования, а не автоматического изменения SEO.")
    return c


def queue(p, summary, analysis):
    c = Canvas(p, "Практический результат:\nготовый список запросов для товаров", "4 · РЕЗУЛЬТАТЫ / ПРАКТИЧЕСКАЯ ВЫДАЧА")
    q = summary["current_queue"]
    c.text(number(q["recommended_review_pairs"]), .6, 3.03, 8.7, 1.35, 80, DARK_MINT, bold=True)
    c.text(f"предложений для {q['review_skus']} товаров", .6, 4.65, 8.7, .8, 30)
    c.text(f"Список: {date.fromisoformat(q['as_of'][:10]):%d.%m.%Y}.\nТрафик по {date.fromisoformat(q['traffic_data_through']):%d.%m.%Y};\nтексты от {date.fromisoformat(q['text_snapshot_day']):%d.%m.%Y}.", .6, 5.9, 8.7, 1.8, 26)
    c.text("Для предложения доступны продающий\nориентир, приоритет и причины отбора.\nСписок открывается в Excel.", .6, 8.1, 8.7, 1.2, 22, GREY)
    for i, (title, detail, color) in enumerate([
        ("Фраза уже есть", f"{number(q['decisions']['already_present'])} кандидатов найдены в тексте — не предлагаем повторно.", BLUE),
        ("Явный конфликт", f"{number(q['decisions']['exclude_candidate'])} кандидатов исключены: бренд, объём, материал или назначение.", PINK),
        ("Приоритетный список", "Запросы и основания рекомендации — для выбора SEO-специалистом.", DARK_MINT)]):
        y = 3.17 + i * 2.04
        c.shape(10.3, y, .1, 1.60, color)
        c.label(title, 10.63, y + .08, 8.77, color)
        c.text(detail, 10.63, y + .73, 8.77, 1.0, 24)
    c.takeaway("Нет фразы в доступных полях ≠ доказано отсутствие во всех SEO-полях.")
    c.note("Источник: experiment_summary.json.current_queue, reports/stage3_release. После отдельного полного обучения сформированы3530review/353SKU, до10наSKU из212787кандидатов. До ограничения списка184825review,27055явных конфликтов,907найденных фраз. Конфликты:объём14001,бренд12411,материал/подтип356,назначение287. Правила консервативны и ограничены словарём. Отсутствие конфликта не является проверенной релевантностью. Полнословное наличие проверяется в доступном тексте, не во всехSEOполях. Score — приоритет, не вероятность. Сохранены selling_seed/history_views/recommendation_basis. Автоматического add_candidate, изменения карточек или прогноза эффекта нет. Экспертная оценка отложена, она не требуется для запуска сохранённых моделей.")
    return c


def conclusion(p, summary, analysis):
    c = Canvas(p, "Задача решена: система работает,\nобучение улучшило результат", "5 · ВЫВОДЫ И ЗАКЛЮЧЕНИЕ")
    a, q = summary["ablation"], summary["current_queue"]
    control, learned = (a["test"][key]["recall@10"] for key in ("e5_co_sales", "all_four"))
    gain = percent(learned - control).removesuffix("%")
    c.label("Вклад обучения · на одинаковом пуле", .6, 3.02, 18.8, BLACK)
    for x, value, label, color in [
        (.6, percent(control), "E5 + совместные продажи", GREY),
        (7.1, percent(learned), "+ адаптер и LightGCN", DARK_MINT),
        (13.6, f"+{gain} п.п.", "Прирост Recall@10", PURPLE),
    ]:
        c.text(value, x, 3.85, 5.8, 1.02, 52, color, bold=True)
        c.text(label, x, 5.01, 5.8, .65, 22, GREY)
    c.line(.6, 5.95, 19.4, 5.95, LIGHT, 2)
    c.label("Рабочий результат", .6, 6.30, 9.1)
    c.text(f"{number(q['recommended_review_pairs'])} предложений / {q['review_skus']} товаров.\nСохранённые модели выдают список\nбез повторного обучения;\nконтрольные суммы проверяются.", .6, 7.04, 9.1, 2.1, 25)
    c.label("Практическая ценность", 10.3, 6.30, 9.1, PURPLE)
    c.text("Новые запросы для каждого товара.\nПриоритеты вместо ручного перебора.\nОтсев дублей и явных конфликтов.\nГотовый список для SEO-специалиста.", 10.3, 7.04, 9.1, 2.1, 25)
    c.takeaway("Создана и проверена рекомендательная система: поиск → обучение → оценка → выдача.")
    c.note("Источники: experiment_summary.json.ablation и current_queue. Recall@10 — средняя по SKU доля будущих продающих кластеров в первой десятке, не процент релевантных фраз. Вклад адаптера и LightGCN проверен на том же пуле, веса вариантов выбраны на validation. Сентябрь уже просмотрен; bootstrap внутри него не даёт нового независимого доказательства. Новый четырёхдневный интервал разницы включает ноль. Работоспособность подтверждена настоящим обучением, запуском восьми этапов и побайтовым воспроизведением выдачи без обучения. Это не равнозначно подтверждению качества на всём каталоге. Отсутствие заказа не означает нерелевантность. Измеренный прирост продаж не заявляется. A/B и рост продаж — последующая бизнес-задача.")
    return c


def development(p, summary, analysis):
    c = Canvas(p, "Как улучшать систему:\nкаждый шаг связан с найденной проблемой", "5 · ВЫВОДЫ / НАПРАВЛЕНИЯ РАЗВИТИЯ")
    distribution = summary["stage4"]["test"]["rank_distribution"]
    below = sum(distribution[key]["macro_share"] for key in ("rank_11_50", "rank_over_50"))
    missing = distribution["missing_from_pool"]["macro_share"]
    c.text("Сначала независимая проверка. Затем — доработка по измеренным причинам ошибок.", .6, 3.00, 18.8, .80, 26)
    for i, (title, reason, action, check, color) in enumerate([
        ("Новый период", "Сентябрь уже использовался;\nоктябрьская проверка — лишь 4 дня.",
         "Зафиксировать модель и проверить\nследующие полные 14 дней.",
         "Recall / NDCG и интервал разницы\nс контрольным гибридом.", BLUE),
        ("Поиск и порядок", f"{percent(missing)} целей не найдены;\n{percent(below)} — ниже первой десятки.",
         "Расширять поиск; проверять пользу\nточного текстового ранжирования.",
         "Полнота пула и NDCG;\nвклад каждого изменения отдельно.", DARK_MINT),
        ("Соответствие товару", "Правила выявляют лишь явные\nконфликты; остальное неизвестно.",
         "Расширять проверку свойств;\nотдельно провести аудит фраз.",
         "Доля конфликтов и подходящих\nфраз; неизвестные — отдельно.", PURPLE),
    ]):
        x = .6 + i * 6.5
        c.shape(x, 4.15, .95, .95, color, oval=True)
        c.text(str(i + 1), x, 4.35, .95, .55, 28, WHITE, bold=True, align="center")
        c.text(title, x, 5.42, 5.8, .70, 28, color, bold=True)
        c.text(reason, x, 6.36, 5.8, .98, 21)
        c.text(action, x, 7.58, 5.8, .98, 21)
        c.text(check, x, 8.80, 5.8, .82, 19, GREY)
    c.takeaway("Фото — вне текущего объёма. Рост заказов проверяется отдельно, после SEO-изменений.")
    c.note("План не объявляет перечисленные эксперименты выполненными. Новый holdout должен начаться после фиксации версии и не участвовать в настройке; 14 дней — текущий горизонт, возможно накопление нескольких окон. Доли потерь — macro по SKU, не доли от260связей. Для поиска: дополнительные каналы и работа с новыми товарами, контроль полноты пула. Для порядка: эксперимент с cross-encoder/дообучением текста только при подтверждённой пользе. Товарный аудит оценивает соответствие подтверждённым свойствам, не угадывает продажи; неизвестные не превращаются в негативы. Экспертная разметка отложена и не блокирует запуск. GraphSAGE, fine-tuning E5, фото/OCR — возможные расширения, не обязательные текущие компоненты. A/B нужен для бизнес-эффекта, не для доказательства запуска модели.")
    return c


def code(p, summary, analysis):
    c = Canvas(p, "Рабочий код и воспроизведение", "МАТЕРИАЛЫ ДЛЯ ЗАЩИТЫ")
    c.text("Репозиторий проекта", .6, 3.02, 18.8, .55, 30, bold=True)
    c.text(REPOSITORY, .6, 3.8, 18.8, .70, 28, BLUE, url=REPOSITORY)
    c.label("Одна точка входа", .6, 5.12, 9.1)
    c.text("src/recommender.py\nrun — весь процесс\nverify — проверка результатов\nrecommend — готовые модели", .6, 5.88, 9.1, 2.50, 26)
    c.label("Материалы рядом с кодом", 10.3, 5.12, 9.1, PURPLE)
    c.text("README.md — запуск и статус\ndocs/recommender_methodology.md\nreports/experiment_summary.json\nreports/explanatory_note.docx", 10.3, 5.88, 9.1, 2.50, 24)
    c.text("Демонстрация: verify проверяет 8 этапов; recommend воспроизводит выдачу без обучения.\n9 ML-модулей и 2 генератора материалов; данные, сертификаты и .env не публикуются.", .6, 8.96, 18.8, .80, 22, GREY)
    c.takeaway("Перед защитой: опубликовать нужный коммит и проверить доступ преподавателей.")
    c.note("Перед защитой предоставить доступ преподавателям и зафиксировать коммит; ссылка не доказывает публикацию локальной версии. Реальные сохранённые модели побайтно воспроизвели обаCSVбезобучения. Команды:.venv/bin/python src/recommender.py verify --config config.json; .venv/bin/python src/recommender.py recommend --release reports/stage3_release --as-of 2026-10-11T01:28:00Z --output reports/recommendations_demo. Путь вывода должен быть новым. Малый тест:.venv/bin/python -m pytest -q tests/test_workflow.py, обучает настоящие модели на синтетических данных; это функциональная проверка, не доказательство качества. Исходные бизнес-данные и веса нужны только для реальной демонстрации, передаются отдельно. Секреты не показывать. README содержит инструкцию сдачи.")
    return c


def sources(p, summary, analysis):
    c = Canvas(p, "Научные источники:\nпочему выбраны эти методы", "6 · СПИСОК ИСТОЧНИКОВ / МЕТОДЫ")
    refs = [("CatBoost · бустинг", "Prokhorenkova et al., 2018.\nCatBoost: unbiased boosting\nwith categorical features.", "https://arxiv.org/abs/1706.09516"),
            ("E5 · смысловые векторы", "Wang et al., 2022. Text Embeddings\nby Weakly-Supervised\nContrastive Pre-training.", "https://arxiv.org/abs/2212.03533"),
            ("LightGCN · граф", "He et al., 2020. LightGCN:\nSimplifying and Powering Graph\nConvolution Network for Recommendation.", "https://arxiv.org/abs/2002.02126"),
            ("BPR · обучение графа", "Rendle et al., UAI 2009.\nBayesian Personalized Ranking\nfrom Implicit Feedback.", "https://arxiv.org/abs/1205.2618"),
            ("RRF · гибрид рангов", "Cormack, Clarke, Büttcher, 2009.\nReciprocal Rank Fusion.\nSIGIR.", "https://doi.org/10.1145/1571941.1572114"),
            ("NDCG · качество порядка", "Järvelin, Kekäläinen, 2002.\nCumulated gain-based evaluation\nof IR techniques. ACM TOIS.", "https://doi.org/10.1145/582415.582418")]
    for i, (title, detail, url) in enumerate(refs):
        x, y = .6 + (i % 2) * 9.7, 3.05 + (i // 2) * 2.17
        c.line(x, y, x + 8.9, y, MINT, 2)
        c.label(title, x, y + .18, 9.1, BLACK)
        c.text(detail, x, y + .83, 9.0, .90, 18, GREY)
        c.text("Статья / DOI →", x, y + 1.81, 9.0, .38, 18, BLUE, url=url)
    c.takeaway("Статьи обосновывают методы. Качество нашей модели измерено отдельными экспериментами.")
    c.note("Полные библиографические названия и активные ссылки — в пояснительной записке. BPR опубликован в UAI2009; запись arXiv1205.2618 размещена2012, это не дата исходной публикации. СтатьяCatBoost обосновывает семейство метода, точнаяPairLogit — в документации на следующем слайде. ИсходнаяE5публикация объясняет подход; использованнаяmultilingual-e5-small документирована отдельно. Не приписываем авторство методов или нормализатораWB себе.")
    return c


def implementation_sources(p, summary, analysis):
    c = Canvas(p, "Источники реализации\nи происхождение результатов", "6 · СПИСОК ИСТОЧНИКОВ / РЕАЛИЗАЦИЯ И ДАННЫЕ")
    refs = [("Sentence Transformers", "Reimers, Gurevych, 2019.\nSentence-BERT · семантический поиск.", "https://arxiv.org/abs/1908.10084"),
            ("multilingual-e5-small", "Карточка intfloat: точная модель,\nразмер вектора и префиксы текстов.", "https://huggingface.co/intfloat/multilingual-e5-small"),
            ("InfoNCE · цель адаптера", "van den Oord, Li, Vinyals, 2018.\nContrastive Predictive Coding.", "https://arxiv.org/abs/1807.03748"),
            ("PairLogit · цель CatBoost", "Документация CatBoost:\nпопарное обучение ранжированию.", "https://catboost.ai/en/docs/concepts/loss-functions-ranking#pairlogit"),
            ("ClickHouse · источник", "Официальная документация:\nчтение и агрегация истории.", "https://clickhouse.com/docs"),
            ("pandas · обработка", "Официальная документация:\nтабличные преобразования.", "https://pandas.pydata.org/docs/")]
    for i, (title, detail, url) in enumerate(refs):
        x, y = .6 + (i % 2) * 9.7, 3.0 + (i // 2) * 1.8
        c.label(title, x, y, 9.0, BLACK)
        c.text(detail, x, y + .59, 9.0, .69, 19, GREY)
        c.text("Открыть источник →", x, y + 1.34, 9.0, .38, 18, BLUE, url=url)
    c.line(.6, 8.47, 19.4, 8.47, LIGHT, 2)
    c.text("Данные: история запросов Ozon + словарь WB + датированный текст карточек.\nНаши числа: experiment_summary.json; входы и результаты связаны контрольными суммами.", .6, 8.77, 18.8, .86, 21)
    c.takeaway("12 внешних источников. Полные ссылки и описание их роли — в пояснительной записке.")
    c.note("История: default_staging.stg_ozon_product_queries_details. Канонические кластеры: raw.wb_normquery, снимок26.06.2026. Текущий текст: raw.ozon_product_info_attributes, снимок06.10, не используется как исторический сентябрьский признак. Происхождение чисел: reports/experiment_summary.json.provenance, stage4.provenance, манифесты оценочных и финальных артефактов. Эти внутренние материалы — источники результатов; внешние статьи их не подтверждают. Карточка модели и документация — изменяемые веб-ресурсы. Дата составления списка11.10.2026; доступность всехссылок в сети не подтверждена. Готовые весаE5заморожены, наш обученный адаптер используетInfoNCE; полнаяархитектураCPC не реализуется. Фото/OCR не используются.")
    return c


BUILDERS = (cover, route, problem, goal, metrics, data_slide, quality, dependencies,
            requirements_slide, pipeline, search, learning, trained_models, temporal, results,
            neural_value, new_period, robustness, ranking_diagnosis, queue, conclusion, development,
            code, sources, implementation_sources)


def validate(canvases):
    if len(canvases) != TOTAL:
        raise ValueError("Неверное число слайдов")
    for index, c in enumerate(canvases, 1):
        for shape in c.slide.shapes:
            if min(shape.left, shape.top) < 0 or shape.left + shape.width > Inches(WIDTH) + 5 or shape.top + shape.height > Inches(HEIGHT) + 5:
                raise ValueError(f"Элемент за границами слайда {index}")
        for i, a in enumerate(c.text_bounds):
            for b in c.text_bounds[i + 1:]:
                if min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]) > .02 and min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]) > .02:
                    raise ValueError(f"Пересечение текста на слайде {index}: {a[4]} / {b[4]}")


def build(summary, analysis):
    if summary["evaluation_status"] != "reused_diagnostic_period":
        raise ValueError("Перед изменением статуса проверки необходимо обновить рассказ")
    if summary["schema_version"] != 2 or summary["selected_method"] != "learned_fusion":
        raise ValueError("Нужен актуальный снимок обученных компонентов")
    if (summary["source"]["fact_rows"] != summary["aggregation"]["input_rows"]
            or summary["analysis_reference"] != {"rows": analysis["source"]["rows"], "date_max": analysis["source"]["date_max"]}
            or summary["source"]["last_day"] < analysis["source"]["date_max"]):
        raise ValueError("Источники анализа и результатов расходятся")
    if summary.get("stage4", {}).get("scope") != "fixed_candidate_pool_reranking":
        raise ValueError("Нужен измеренный снимок части 4; обновите агрегаты с --results")
    p = Presentation()
    p.slide_width, p.slide_height = Inches(WIDTH), Inches(HEIGHT)
    p.core_properties.title = "Рекомендательная система поисковых кластеров Ozon"
    p.core_properties.subject = "Постановка задачи, анализ, методика, результаты, выводы и источники"
    p.core_properties.author = ""
    p.core_properties.last_modified_by = ""
    apply_theme(p)
    canvases = [builder(p, summary, analysis) for builder in BUILDERS]
    validate(canvases)
    return p, canvases


def save_materials(output=DECK, preview=None):
    if output.suffix.lower() != ".pptx" or (preview and preview.suffix.lower() != ".png"):
        raise ValueError("Выходные файлы должны иметь расширения .pptx и .png")
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    p, canvases = build(summary, analysis)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent, prefix="defense-") as folder:
        temporary = Path(folder) / "presentation.pptx"
        p.save(temporary)
        if len(Presentation(temporary).slides) != TOTAL:
            raise ValueError("Сохранение изменило число слайдов")
        temporary.replace(output)
    if preview:
        preview.parent.mkdir(parents=True, exist_ok=True)
        sheet = Image.new("RGB", (1600, math.ceil(TOTAL / 3) * 315), "#E8EAEE")
        for i, canvas in enumerate(canvases):
            tile = canvas.image.resize((520, 293))
            sheet.paste(tile, ((i % 3) * 535, (i // 3) * 315))
        sheet.save(preview)
    return p


def parse_pages(text: str) -> list[list[str]]:
    pages = [re.split(r"\n\s*\n", page.strip()) for page in text.split(PAGEBREAK)]
    if len(pages) != 10:
        raise ValueError("Источник должен содержать ровно десять страниц")
    for number, blocks in enumerate(pages, 1):
        if not blocks[0].startswith("# ") or "\n" in blocks[0]:
            raise ValueError(f"Страница {number} должна начинаться с заголовка")
        if any(block.startswith("#") for block in blocks[1:]):
            raise ValueError(f"Лишний заголовок на странице {number}")
        if len(" ".join(blocks).split()) > 400:
            raise ValueError(f"Страница {number} превышает бюджет 400 слов")
    return pages


def add_link(paragraph, label: str, target: str, base_path=None, anchors=None) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    if target.startswith("#"):
        if anchors is None or target[1:] not in anchors:
            raise ValueError(f"Ссылка на отсутствующий раздел: {target}")
        hyperlink.set(qn("w:anchor"), anchors[target[1:]])
    else:
        if base_path is not None and not re.match(r"[a-zA-Z][\w+.-]*:", target):
            resolved_target = (base_path / target).resolve()
            if resolved_target.is_relative_to(ROOT):
                target = f"{REPOSITORY}/blob/main/{resolved_target.relative_to(ROOT).as_posix()}"
            else:
                target = resolved_target.as_uri()
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(
            target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True,
        ))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    for tag, value in (("color", "226655"), ("u", "single")):
        element = OxmlElement(f"w:{tag}")
        element.set(qn("w:val"), value)
        properties.append(element)
    run.append(properties)
    content = OxmlElement("w:t")
    content.set(qn("xml:space"), "preserve")
    content.text = label.strip("`")
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_text(paragraph, text: str, base_path=None, anchors=None) -> None:
    position = 0
    for match in INLINE.finditer(text):
        paragraph.add_run(text[position:match.start()])
        token = match[0]
        if link := LINK.fullmatch(token):
            add_link(paragraph, link[1], link[2], base_path, anchors)
        elif token.startswith("`"):
            paragraph.add_run(token[1:-1]).font.name = "Courier New"
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token[1:-1]).italic = True
        position = match.end()
    paragraph.add_run(text[position:])


def split_table_row(line: str) -> list[str]:
    cells, cell, in_code = [], [], False
    position = 0
    line = line.strip().removeprefix("|")
    while position < len(line):
        character = line[position]
        if character == "\\" and position + 1 < len(line) and line[position + 1] == "|":
            cell.append("|")
            position += 2
            continue
        if character == "`":
            in_code = not in_code
        if character == "|" and not in_code:
            cells.append("".join(cell).strip())
            cell = []
        else:
            cell.append(character)
        position += 1
    if cell or not line.endswith("|"):
        cells.append("".join(cell).strip())
    return cells


def add_table(document, block: str, base_path=None, anchors=None) -> None:
    rows = [split_table_row(line) for line in block.splitlines()]
    if len(rows) < 3 or not all(re.fullmatch(r":?-+:?", cell) for cell in rows[1]):
        raise ValueError("Неверная Markdown-таблица")
    rows.pop(1)
    columns = len(rows[0])
    if any(len(row) != columns for row in rows):
        raise ValueError("Неодинаковое число столбцов таблицы")
    widths = {2: (55, 115), 3: (53, 47, 70), 4: (65, 30, 43, 32)}.get(columns)
    if widths is None:
        if base_path is None:
            raise ValueError("Поддерживаются таблицы с 2–4 столбцами")
        widths = (170 / columns,) * columns
    if rows[0][-1] == "Страница":
        widths = (145, 25)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style, table.autofit = "Table Grid", False
    for column, width in zip(table.columns, widths):
        column.width = docx.shared.Mm(width)
    for index, (row, values) in enumerate(zip(table.rows, rows)):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = docx.shared.Mm(width)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = docx.shared.Pt(2)
            paragraph.paragraph_format.space_before = docx.shared.Pt(2)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_text(paragraph, value, base_path, anchors)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EAF5F0")
                cell._tc.get_or_add_tcPr().append(shading)
    spacing = document.add_paragraph()
    spacing.paragraph_format.space_after = docx.shared.Pt(0)
    spacing.paragraph_format.line_spacing = docx.shared.Pt(3)
    spacing.add_run().font.size = docx.shared.Pt(3)


def word_document(guide=False):
    """Общие настройки двух документов: A4, поля, стили и нумерация."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = docx.shared.Mm(210), docx.shared.Mm(297)
    section.top_margin = section.bottom_margin = docx.shared.Mm(20)
    section.left_margin = section.right_margin = docx.shared.Mm(20)
    section.header_distance = section.footer_distance = docx.shared.Mm(8)
    sizes = (("Normal", 11), ("Title", 20), ("Heading 1", 15), ("Heading 2", 13),
             ("Heading 3", 12)) if guide else (("Normal", 12), ("Title", 12), ("Heading 1", 12))
    for name, size in sizes:
        style = document.styles[name]
        style.font.name, style.font.size = "Times New Roman", docx.shared.Pt(size)
        style.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            style.element.get_or_add_rPr().rFonts.set(qn(f"w:{script}"), "Times New Roman")
        style.paragraph_format.line_spacing = 1.1 if guide else 1.15
        style.paragraph_format.space_before = docx.shared.Pt(0)
        style.paragraph_format.space_after = docx.shared.Pt(6)
        style.paragraph_format.widow_control = True
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
    document.styles["Title"].font.bold = True
    if guide:
        code_style = document.styles.add_style("Guide Code", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name, code_style.font.size = "Courier New", docx.shared.Pt(9)
        code_style.paragraph_format.line_spacing = 1
        code_style.paragraph_format.space_before = docx.shared.Pt(4)
        code_style.paragraph_format.space_after = docx.shared.Pt(8)
        code_style.paragraph_format.widow_control = True
    else:
        document.styles["Heading 1"].paragraph_format.space_after = docx.shared.Pt(12)
    header = section.header.paragraphs[0]
    header.style = document.styles["Normal"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("Личный гайд по рекомендательной системе" if guide else "Малотрафиковые поисковые кластеры")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for instruction in (("PAGE", "NUMPAGES") if guide else ("PAGE",)):
        if instruction == "NUMPAGES":
            footer.add_run(" / ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction)
        footer._p.append(field)
    document.core_properties.author = "[ФИО автора]"
    document.core_properties.last_modified_by = ""
    document.core_properties.created = document.core_properties.modified = datetime(2026, 10, 11, tzinfo=timezone.utc)
    return document


def create_explanatory_note(source: Path = DEFAULT_SOURCE, output: Path = DEFAULT_OUTPUT) -> Path:
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve() or output.suffix.lower() != ".docx":
        raise ValueError("Выход должен быть отдельным файлом .docx")
    pages = parse_pages(source.read_text(encoding="utf-8"))
    document = word_document()
    for number, blocks in enumerate(pages, 1):
        if number > 1:
            page_break = document.add_page_break()
            page_break.paragraph_format.space_before = docx.shared.Pt(0)
            page_break.paragraph_format.space_after = docx.shared.Pt(0)
            page_break.paragraph_format.line_spacing = docx.shared.Pt(1)
            page_break.runs[0].font.size = docx.shared.Pt(1)
        heading = document.add_paragraph(blocks[0][2:], "Title" if number == 1 else "Heading 1")
        if number == 1:
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            heading.paragraph_format.space_before = docx.shared.Pt(50)
            heading.paragraph_format.space_after = docx.shared.Pt(24)
        for index, block in enumerate(blocks[1:]):
            if block.startswith("|"):
                add_table(document, block)
            else:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if number == 1 and index < 7:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_after = docx.shared.Pt(18)
                add_text(paragraph, " ".join(block.splitlines()))
    document.core_properties.title = pages[0][0][2:]
    document.core_properties.subject = "Дипломный проект: обученная гибридная рекомендательная система"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


def parse_guide_blocks(text: str) -> list[tuple[str, str]]:
    lines, blocks, position = text.splitlines(), [], 0
    while position < len(lines):
        line = lines[position].strip()
        if not line:
            position += 1
            continue
        fence = re.fullmatch(r"(`{3,}|~{3,})(.*)", line)
        if fence:
            code = []
            position += 1
            while position < len(lines) and not re.fullmatch(
                re.escape(fence[1][0]) + "{" + str(len(fence[1])) + r",}\s*", lines[position].strip(),
            ):
                code.append(lines[position])
                position += 1
            if position == len(lines):
                raise ValueError("Незакрытый блок кода")
            blocks.append(("code", "\n".join(code)))
            position += 1
        elif heading := re.fullmatch(r"(#{1,6})\s+(.+)", line):
            blocks.append((f"heading_{len(heading[1])}", heading[2]))
            position += 1
        elif line.startswith("|"):
            table = []
            while position < len(lines) and lines[position].strip().startswith("|"):
                table.append(lines[position].strip())
                position += 1
            blocks.append(("table", "\n".join(table)))
        elif item := re.fullmatch(r"(?:[-+*]|\d+\.)\s+(.+)", line):
            blocks.append(("bullet" if line[0] in "-+*" else "number", item[1]))
            position += 1
        else:
            paragraph = [line]
            position += 1
            while position < len(lines) and lines[position].strip() and not re.match(
                r"^(?:#{1,6}\s|`{3,}|~{3,}|\||[-+*]\s|\d+\.\s)", lines[position].strip(),
            ):
                paragraph.append(lines[position].strip())
                position += 1
            blocks.append(("paragraph", " ".join(paragraph)))
    if not blocks or blocks[0][0] != "heading_1":
        raise ValueError("Гайд должен начинаться с заголовка первого уровня")
    return blocks


def create_project_guide(source: Path = DEFAULT_GUIDE_SOURCE, output: Path = DEFAULT_GUIDE_OUTPUT) -> Path:
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve() or output.suffix.lower() != ".docx":
        raise ValueError("Выход должен быть отдельным файлом .docx")
    blocks = parse_guide_blocks(source.read_text(encoding="utf-8"))
    document = word_document(guide=True)
    heading_links, anchors, occurrences = {}, {}, {}
    for index, (kind, value) in enumerate(blocks[1:], 1):
        if kind.startswith("heading_"):
            slug = re.sub(r"[^\w\s-]", "", value.lower()).replace(" ", "-")
            occurrence = occurrences.get(slug, 0)
            occurrences[slug] = occurrence + 1
            slug = f"{slug}-{occurrence}" if occurrence else slug
            anchors[slug] = f"section_{index}"
            heading_links[index] = slug
    add_text(document.add_paragraph(style="Title"), blocks[0][1])
    document.add_paragraph("Содержание", "Heading 1")
    for index, slug in heading_links.items():
        kind, value = blocks[index]
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = docx.shared.Mm(max(0, int(kind[-1]) - 2) * 5)
        paragraph.paragraph_format.space_after = docx.shared.Pt(2)
        add_link(paragraph, value, f"#{slug}", source.parent, anchors)
    for index, (kind, value) in enumerate(blocks[1:], 1):
        if kind.startswith("heading_"):
            paragraph = document.add_paragraph(style=f"Heading {max(1, int(kind[-1]) - 1)}")
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(index))
            start.set(qn("w:name"), anchors[heading_links[index]])
            paragraph._p.append(start)
            add_text(paragraph, value, source.parent, anchors)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(index))
            paragraph._p.append(end)
        elif kind == "table":
            add_table(document, value, source.parent, anchors)
        elif kind == "code":
            document.add_paragraph(value, "Guide Code")
        else:
            style = {"bullet": "List Bullet", "number": "List Number"}.get(kind, "Normal")
            add_text(document.add_paragraph(style=style), value, source.parent, anchors)
    document.core_properties.title = blocks[0][1]
    document.core_properties.subject = "Понятная методика, воспроизведение и подготовка к защите"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


def main(argv: list[str] | None = None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    presentation = commands.add_parser("presentation", help="Пересобрать PowerPoint из публичных агрегатов")
    presentation.add_argument("--output", type=Path, default=DECK)
    presentation.add_argument("--preview", type=Path)
    for name in ("note", "guide"):
        command = commands.add_parser(name, help="Собрать записку" if name == "note" else "Собрать личный гайд")
        command.add_argument("--source", type=Path)
        command.add_argument("--output", type=Path)
    summary = commands.add_parser("summary", help="Обновить агрегаты из результатов, без обучения")
    summary.add_argument("--config", type=Path, default=ROOT / "config.json")
    summary.add_argument("--output", type=Path, default=SUMMARY)
    summary.add_argument("--results", type=Path, default=ROOT / "reports/stage4_results/metrics.json")
    args = parser.parse_args(argv)
    if args.command == "presentation":
        save_materials(args.output, args.preview)
        print(f"{TOTAL} слайдов: {args.output}")
    elif args.command == "summary":
        refresh_summary(args.config, args.output, args.results)
        print(f"Обновлены агрегаты: {args.output}")
    elif args.command == "guide":
        result = create_project_guide(args.source or DEFAULT_GUIDE_SOURCE, args.output or DEFAULT_GUIDE_OUTPUT)
        print(f"Создан гайд: {result}; пагинацию проверьте в редакторе")
    else:
        result = create_explanatory_note(args.source or DEFAULT_SOURCE, args.output or DEFAULT_OUTPUT)
        print(f"Создана записка: {result}; разделов: 10; пагинацию проверьте в редакторе")


if __name__ == "__main__":
    main()